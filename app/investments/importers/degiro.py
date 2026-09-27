"""
app/investments/importers/degiro.py
===================================
FAMILYFIN — Read a DEGIRO "Estado de cuenta" (Account statement) export.

The statement lists every cash movement, newest first, in 12 columns:
  Fecha | Hora | Fecha valor | Producto | ISIN | Descripción | Tipo (FX rate) |
  Variación (currency) | (amount) | Saldo (currency) | (amount) | ID Orden
(.xlsx or .csv, Spanish or English headers/descriptions.)

What becomes an operation
  * "Compra/Buy N <name>@<price> <CUR> (<ISIN>)" rows. Partial executions of
    one order share its "ID Orden" and are merged into one operation.
  * Rows with the same order id that are not the trade itself are its costs
    (DEGIRO fee, Spanish/French/Italian transaction tax) → fees, added to the
    purchase cost / deducted from the sale proceeds. For a trade in another
    currency, AutoFX rows with the same order id give the EUR amount.
  * "Dividendo/Dividend" and "Retención del dividendo/Dividend Tax" rows of a
    product on the same day form one dividend. DEGIRO converts foreign cash
    to EUR in a later "Retirada/Ingreso Cambio de Divisa" pair: the EUR really
    credited is spread over the dividends that conversion covered.
Everything else (deposits, cash sweeps, interest, yearly exchange
connectivity fees) is listed as ignored.

Nothing here touches the database or the network.
"""
from __future__ import annotations

import csv
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional

BROKER = "DEGIRO"

_TRADE = re.compile(
    r"^(Compra|Venta|Buy|Sell) ([\d.,]+) (.+)@([\d.,]+) ([A-Z]{3}) \(([A-Z0-9]{12})\)$")
_DIVIDEND = ("Dividendo", "Dividend")
_DIVIDEND_TAX = ("Retención del dividendo", "Dividend Tax")
_FX_OUT = ("Retirada Cambio de Divisa", "FX Withdrawal", "FX Debit")
_FX_IN = ("Ingreso Cambio de Divisa", "FX Credit")
_HEADERS = ("Fecha", "Date")


class StatementError(ValueError):
    """The file is not a DEGIRO account statement. Message in Spanish."""


@dataclass
class ParsedTrade:
    ref:        str                 # DEGIRO order id
    kind:       str                 # "buy" | "sell"
    trade_date: date
    isin:       str
    name:       str
    quantity:   float
    price:      float               # average of the executions, trade currency
    currency:   str
    amount_eur: Optional[float]     # buy: total charged incl. fees; sell: net received
    fees_eur:   float = 0.0
    fills:      int = 1
    product:    str = ""            # DEGIRO's short product name (column "Producto")


@dataclass
class ParsedDividend:
    ref:          str
    pay_date:     date
    isin:         str
    name:         str
    currency:     str
    gross:        float             # in `currency`
    withheld:     float             # in `currency`, positive
    net_eur:      Optional[float] = None
    withheld_eur: Optional[float] = None
    rate:         Optional[float] = None    # units of currency per EUR used


@dataclass
class ParsedStatement:
    trades:    List[ParsedTrade] = field(default_factory=list)
    dividends: List[ParsedDividend] = field(default_factory=list)
    ignored:   Dict[str, int] = field(default_factory=dict)      # description → rows
    warnings:  List[str] = field(default_factory=list)
    first_day: Optional[date] = None
    last_day:  Optional[date] = None
    broker:    str = BROKER

    @property
    def products(self) -> Dict[str, dict]:
        """ISIN → {name, currency, trades: [(date, price)]} for symbol lookup."""
        result: Dict[str, dict] = {}
        for trade in self.trades:
            item = result.setdefault(trade.isin, {"name": trade.name, "currency": trade.currency,
                                                  "product": trade.product, "trades": []})
            item["trades"].append((trade.trade_date, trade.price))
        for dividend in self.dividends:
            # A dividend's currency is the payout's, not the listing's: assume EUR
            # listings like the trades (the user can change it in the preview).
            result.setdefault(dividend.isin, {"name": dividend.name, "currency": "EUR",
                                              "product": dividend.name, "trades": []})
        return result


# ---------------------------------------------------------------------------
# Reading rows
# ---------------------------------------------------------------------------
def _number(value) -> Optional[float]:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "")
    if "," in text:                       # Spanish: 1.234,56
        text = text.replace(".", "").replace(",", ".")
    return float(text)


def _read_rows(path: Path) -> List[list]:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return [list(r) for r in workbook.worksheets[0].iter_rows(values_only=True)]
        finally:
            workbook.close()
    if suffix == ".csv":
        text = path.read_text(encoding="utf-8-sig")
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;")
        return [row for row in csv.reader(text.splitlines(), dialect)]
    raise StatementError("Formato no admitido: usa el .xlsx o .csv exportado de DEGIRO.")


@dataclass
class _Row:
    when:     datetime
    product:  str
    isin:     str
    desc:     str
    rate:     Optional[float]
    currency: str
    amount:   Optional[float]
    order_id: str


def _parse_rows(raw: List[list]) -> List[_Row]:
    if not raw or not raw[0] or str(raw[0][0]).strip() not in _HEADERS or len(raw[0]) < 12:
        raise StatementError(
            "El archivo no parece un «Estado de cuenta» de DEGIRO (faltan las columnas "
            "Fecha, Hora, Producto, ISIN, Descripción, Variación e ID Orden).")
    rows = []
    for values in raw[1:]:
        values = list(values) + [None] * (12 - len(values))
        if not values[0]:
            continue
        day, clock = str(values[0]).strip(), str(values[1] or "00:00").strip()
        try:
            when = datetime.strptime(f"{day} {clock}", "%d-%m-%Y %H:%M")
        except ValueError as exc:
            raise StatementError(f"Fecha no válida en el extracto: {day} {clock}") from exc
        rows.append(_Row(
            when=when,
            product=str(values[3] or "").strip(),
            isin=str(values[4] or "").strip(),
            desc=str(values[5] or "").strip(),
            rate=_number(values[6]),
            currency=str(values[7] or "").strip(),
            amount=_number(values[8]),
            order_id=str(values[11] or "").strip(),
        ))
    rows.sort(key=lambda r: r.when)       # the export is newest first
    return rows


# ---------------------------------------------------------------------------
# Interpretation
# ---------------------------------------------------------------------------
def parse_statement(path) -> ParsedStatement:
    rows = _parse_rows(_read_rows(Path(path)))
    statement = ParsedStatement()
    if rows:
        statement.first_day, statement.last_day = rows[0].when.date(), rows[-1].when.date()

    orders: Dict[str, dict] = {}
    order_rows: Dict[str, List[_Row]] = defaultdict(list)
    dividend_rows: List[_Row] = []
    conversions: Dict[datetime, dict] = {}

    for row in rows:
        match = _TRADE.match(row.desc)
        if match and row.order_id:
            order = orders.setdefault(row.order_id, {
                "kind": "buy" if match.group(1) in ("Compra", "Buy") else "sell",
                "date": row.when.date(), "name": match.group(3).strip(),
                "currency": match.group(5), "isin": match.group(6),
                "qty": 0.0, "cash": 0.0, "fills": 0, "cash_currency": row.currency,
                "product": row.product,
            })
            order["qty"] += _number(match.group(2))
            order["cash"] += row.amount or 0.0
            order["fills"] += 1
        elif row.order_id:
            order_rows[row.order_id].append(row)
        elif row.desc in _DIVIDEND or row.desc in _DIVIDEND_TAX:
            dividend_rows.append(row)
        elif row.desc in _FX_OUT or row.desc in _FX_IN:
            conv = conversions.setdefault(row.when, {"when": row.when, "foreign": 0.0,
                                                     "currency": None, "eur": 0.0, "rate": None})
            if row.currency == "EUR":
                conv["eur"] += row.amount or 0.0
            else:
                conv["currency"], conv["foreign"] = row.currency, -(row.amount or 0.0)
                conv["rate"] = row.rate
        else:
            label = re.sub(r"[\d.,]+\s*EUR$|\s*\(.*\)$|:.*$", "", row.desc).strip() or "Otros"
            statement.ignored[label] = statement.ignored.get(label, 0) + 1

    for ref, order in orders.items():
        statement.trades.append(_build_trade(ref, order, order_rows.get(ref, []), statement))
    statement.trades = [t for t in statement.trades if t is not None]
    statement.trades.sort(key=lambda t: (t.trade_date, t.ref))
    statement.dividends = _build_dividends(dividend_rows, conversions, statement)
    # Dividends only carry DEGIRO's upper-case product name: reuse the trade's.
    names = {t.isin: t.name for t in statement.trades}
    for dividend in statement.dividends:
        dividend.name = names.get(dividend.isin, dividend.name)
    return statement


def _build_trade(ref, order, extra: List[_Row], statement) -> Optional[ParsedTrade]:
    fees = -sum(r.amount or 0.0 for r in extra
                if r.currency == "EUR" and r.desc not in _FX_OUT + _FX_IN)
    quantity = order["qty"]
    price = abs(order["cash"]) / quantity if quantity else 0.0
    if order["cash_currency"] == "EUR":
        gross_eur = abs(order["cash"])
    else:
        # AutoFX: the EUR side of the conversion rows booked with this order.
        eur_side = [r.amount or 0.0 for r in extra if r.currency == "EUR"
                    and r.desc in _FX_OUT + _FX_IN]
        if not eur_side:
            statement.warnings.append(
                f"{order['name']} ({order['date']:%d/%m/%Y}): operación en {order['cash_currency']} "
                "sin el cambio a EUR en el extracto; no se importa.")
            return None
        gross_eur = abs(sum(eur_side))
    amount = gross_eur + fees if order["kind"] == "buy" else gross_eur - fees
    return ParsedTrade(ref=ref, kind=order["kind"], trade_date=order["date"], isin=order["isin"],
                       name=order["name"], quantity=quantity, price=round(price, 6),
                       currency=order["currency"], amount_eur=round(amount, 2),
                       fees_eur=round(fees, 2), fills=order["fills"], product=order["product"])


def _build_dividends(rows: List[_Row], conversions: Dict[datetime, dict],
                     statement) -> List[ParsedDividend]:
    grouped: Dict[tuple, ParsedDividend] = {}
    first_seen: Dict[tuple, datetime] = {}
    for row in rows:
        key = (row.isin, row.when.date())
        item = grouped.get(key)
        if item is None:
            # DEGIRO sometimes books the tax a day apart from the dividend.
            near = [k for k in grouped if k[0] == row.isin
                    and abs((k[1] - row.when.date()).days) <= 3]
            key = near[0] if near else key
            item = grouped.get(key)
        if item is None:
            item = grouped[key] = ParsedDividend(
                ref=f"DIV|{row.isin}|{row.when.date().isoformat()}", pay_date=row.when.date(),
                isin=row.isin, name=row.product, currency=row.currency, gross=0.0, withheld=0.0)
            first_seen[key] = row.when
        first_seen[key] = max(first_seen[key], row.when)
        if row.desc in _DIVIDEND:
            item.gross += row.amount or 0.0
        else:
            item.withheld += -(row.amount or 0.0)

    dividends = [d for d in grouped.values() if d.gross - d.withheld > 0.004]
    pending: Dict[str, List[ParsedDividend]] = defaultdict(list)
    keys = {id(d): k for k, d in grouped.items()}
    events = sorted([(first_seen[keys[id(d)]], 0, d) for d in dividends]
                    + [(c["when"], 1, c) for c in conversions.values()],
                    key=lambda e: (e[0], e[1]))
    for _when, kind, payload in events:
        if kind == 0:
            if payload.currency == "EUR":
                payload.net_eur = round(payload.gross - payload.withheld, 2)
                payload.withheld_eur = round(payload.withheld, 2)
                payload.rate = 1.0
            else:
                pending[payload.currency].append(payload)
            continue
        conv = payload
        queue = pending.get(conv["currency"] or "", [])
        foreign_left, eur_left = conv["foreign"], conv["eur"]
        while queue and foreign_left > 0.004:
            dividend = queue[0]
            net = dividend.gross - dividend.withheld
            if net > foreign_left + 0.01:
                break
            share = eur_left * net / foreign_left if foreign_left else 0.0
            dividend.net_eur = round(share, 2)
            dividend.rate = conv["rate"] or (net / share if share else None)
            dividend.withheld_eur = round(dividend.withheld / dividend.rate, 2) if dividend.rate else None
            foreign_left -= net
            eur_left -= share
            queue.pop(0)

    for queue in pending.values():
        for dividend in queue:
            statement.warnings.append(
                f"Dividendo de {dividend.name} del {dividend.pay_date:%d/%m/%Y}: no se encontró su "
                f"cambio a EUR en el extracto; se estimará con el tipo de cambio de ese día.")
    return sorted(dividends, key=lambda d: (d.pay_date, d.name))
