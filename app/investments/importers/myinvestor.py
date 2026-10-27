"""
app/investments/importers/myinvestor.py
=======================================
FAMILYFIN — Read MyInvestor's account movements export.

Columns (";" separated, decimal comma, newest first):
    Fecha de operación ; Fecha de valor ; Concepto ; Importe ; Divisa

The concept is the only description, truncated to ~30 characters:
  * "NVIDIA CORP @ 7"                 stock/ETF movement, 7 = shares
  * "LABORATORIOS FARMACEUTICOS ROV"  stock movement without the shares
  * "FIDELITY MSCI JAPAN INDEX P AC"  fund subscription/redemption (units unknown)
  * "PERIODO 09/08/2026 09/09/2026"   account interest            → ignored
  * "maria", "Transferencia …"        money in/out of the account → ignored
A negative amount is a purchase. A positive one is a sale, unless it is
small against what was invested in that product (< DIVIDEND_SHARE), in
which case it is a (net) dividend. Both guesses, and every estimated
quantity, can be changed in the import preview.

Missing quantities are estimated once the user has chosen each product's
symbol: amount ÷ that day's close (fund NAV). Shares are rounded to whole
numbers when the gap is small (the amount includes the commission), and a
sale close to everything held is taken as a sale of everything held.
"""
from __future__ import annotations

import csv
import difflib
import re
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from app.investments.importers.base import (
    KIND_BUY,
    KIND_DIVIDEND,
    KIND_IGNORE,
    KIND_SELL,
    ImportRow,
    Statement,
)
from app.investments.importers.degiro import StatementError

BROKER = "MyInvestor"
HEADER = ("Fecha de operación", "Fecha de valor", "Concepto", "Importe", "Divisa")

DIVIDEND_SHARE = 0.05        # a credit below 5 % of what is invested is a dividend
SELL_ALL_SHARE = 0.95        # a sale above 95 % of what is held sells everything
ROUND_SHARES_GAP = 0.03      # shares within 3 % of a whole number are rounded

_INTEREST = re.compile(r"^PERIODO \d\d/\d\d/\d{4} \d\d/\d\d/\d{4}$")
_QUANTITY = re.compile(r"^(?P<name>.*?)\s*@\s*(?P<qty>[\d.,]*)\s*$")
_FUND_WORDS = re.compile(r"\b(INDEX|INDX|FUND|FONDO|FI|IDX)\b")


def looks_like_myinvestor(path) -> bool:
    try:
        first = _read(Path(path))[0]
    except (StatementError, OSError, IndexError):
        return False
    return [str(c).strip() for c in first[:5]] == list(HEADER)


def _read(path: Path) -> List[list]:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            sheet.reset_dimensions()       # a wrong <dimension> would hide rows (see degiro.py)
            return [list(r) for r in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
    if suffix == ".csv":
        text = path.read_text(encoding="utf-8-sig")
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=";,")
        return [row for row in csv.reader(text.splitlines(), dialect) if row]
    raise StatementError("Formato no admitido: usa el .csv o .xlsx exportado de MyInvestor.")


def _amount(value) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    return float(text or 0)


def _day(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.strptime(str(value).strip(), "%d/%m/%Y").date()


def product_key(name: str) -> str:
    return "NAME:" + re.sub(r"[^A-Z0-9]", "", name.upper())


# Fund names are abbreviated freely ("ISHARES DEVL WRLD INDX D EUR" for the
# purchase, "ISHARES DEVELOPED WORLD INDEX" for the redemption): compare them
# with the abbreviations unified.
_FUND_WORD_FORMS = {"DEVELOPED": "DEV", "DEVL": "DEV", "WORLD": "WLD", "WRLD": "WLD",
                    "INDEX": "IDX", "INDX": "IDX", "EMERGING": "EM", "EMERG": "EM",
                    "MARKETS": "MKT", "MKTS": "MKT"}


def _fund_form(name: str) -> str:
    words = re.sub(r"[^A-Z0-9 ]", " ", name.upper()).split()
    return "".join(_FUND_WORD_FORMS.get(w, w) for w in words)


def _same_product(a: str, b: str, fund_names: tuple = ()) -> bool:
    """Two truncated spellings of one security ("…FARMACEUTICOS ROV" / "…FARMACEUTIC ROVI").

    For two funds (`fund_names` = their original names) the same first ten
    letters of the unified form are enough.
    """
    if fund_names:
        fa, fb = (_fund_form(n) for n in fund_names)
        if fa[:10] == fb[:10]:
            return True
    a, b = a[5:], b[5:]
    return a[:12] == b[:12] and difflib.SequenceMatcher(None, a, b).ratio() >= 0.85


def is_fund_name(name: str) -> bool:
    return bool(_FUND_WORDS.search(name.upper()))


def parse_statement(path) -> Statement:
    raw = _read(Path(path))
    if not raw or [str(c).strip() for c in raw[0][:5]] != list(HEADER):
        raise StatementError(
            "El archivo no parece un extracto de movimientos de MyInvestor (faltan las columnas "
            "Fecha de operación, Fecha de valor, Concepto, Importe y Divisa).")
    statement = Statement(broker=BROKER)
    entries = []
    for values in raw[1:]:
        values = list(values) + [None] * (5 - len(values))
        if not values[0]:
            continue
        try:
            entries.append((_day(values[0]), _day(values[1] or values[0]),
                            str(values[2] or "").strip(), _amount(values[3]),
                            str(values[4] or "EUR").strip() or "EUR"))
        except ValueError as exc:
            raise StatementError(f"Fila no válida en el extracto: {values}") from exc
    entries.reverse()                     # oldest first (the export is newest first)
    if entries:
        statement.first_day = min(e[0] for e in entries)
        statement.last_day = max(e[0] for e in entries)

    def ignore(label: str) -> None:
        statement.ignored[label] = statement.ignored.get(label, 0) + 1

    keys: List[str] = []
    invested: Dict[str, float] = defaultdict(float)
    seen_refs: Dict[str, int] = defaultdict(int)
    for op_day, _value_day, concept, amount, currency in entries:
        if not concept or abs(amount) < 0.005:
            ignore("Movimientos sin importe")
            continue
        if _INTEREST.match(concept):
            ignore("Intereses de la cuenta")
            continue
        if concept.lower().startswith("transferencia") or concept != concept.upper():
            # Securities come in upper case; names and transfers are money in/out.
            ignore("Ingresos y transferencias de dinero")
            continue
        match = _QUANTITY.match(concept)
        name = (match.group("name") if match else concept).strip()
        quantity = None
        if match and match.group("qty"):
            quantity = _amount(match.group("qty"))
        fund = is_fund_name(name)
        key = product_key(name)
        key = next((k for k in keys if k == key or _same_product(
            k, key, (statement.products[k]["name"], name)
            if fund and statement.products[k]["fund"] else ())), key)
        if key not in keys:
            keys.append(key)
        product = statement.products.setdefault(key, {
            "name": name, "currency": currency, "product": name, "trades": [], "fund": fund,
            "spellings": set()})
        product["spellings"].add(name)

        if amount < 0:
            kind = KIND_BUY
            invested[key] += -amount
        elif invested[key] > 0 and amount < DIVIDEND_SHARE * invested[key]:
            kind = KIND_DIVIDEND
        elif invested[key] > 0:
            kind = KIND_SELL
            invested[key] = max(0.0, invested[key] - amount)
        else:
            kind = KIND_IGNORE
        base_ref = f"MYI|{op_day.isoformat()}|{key[5:]}|{amount:.2f}"
        seen_refs[base_ref] += 1
        ref = base_ref if seen_refs[base_ref] == 1 else f"{base_ref}#{seen_refs[base_ref]}"
        row = ImportRow(
            ref=ref, day=op_day, key=key, name=name, kind=kind, amount_eur=round(abs(amount), 2),
            quantity=quantity if kind in (KIND_BUY, KIND_SELL) else None, currency=currency,
            kind_editable=True,
            quantity_estimated=kind in (KIND_BUY, KIND_SELL) and quantity is None,
            note=("sin compras previas: revisa el tipo" if kind == KIND_IGNORE else
                  "dividendo neto (sin retención en el extracto)" if kind == KIND_DIVIDEND else
                  "participaciones estimadas con el valor liquidativo" if fund and quantity is None
                  else "acciones estimadas con el precio del día" if quantity is None else ""),
        )
        if row.quantity:
            row.price = round(row.amount_eur / row.quantity, 4)
            product["trades"].append((op_day, row.price))
        statement.rows.append(row)
    for product in statement.products.values():
        product["spellings"] = sorted(product["spellings"])
    return statement


def estimate_quantities(
    rows: List[ImportRow],
    symbol_of: Dict[str, str],
    closes: Dict[str, List[Tuple[date, float]]],
    funds: Dict[str, bool],
) -> None:
    """Fill the estimated quantities in place (chronological, per symbol)."""
    held: Dict[str, float] = defaultdict(float)
    for row in sorted(rows, key=lambda r: (r.day, r.ref)):
        symbol = symbol_of.get(row.key)
        if row.kind not in (KIND_BUY, KIND_SELL) or not symbol:
            continue
        if row.quantity_estimated:
            series = closes.get(symbol, [])
            before = [c for d, c in series if d <= row.day and (row.day - d).days <= 7]
            after = [c for d, c in series if d > row.day and (d - row.day).days <= 5]
            close = before[-1] if before else (after[0] if after else None)
            if close:
                quantity = row.amount_eur / close
                if not funds.get(row.key):
                    whole = round(quantity)
                    if whole and abs(quantity - whole) <= ROUND_SHARES_GAP * whole:
                        quantity = float(whole)
                if row.kind == KIND_SELL and held[symbol] and quantity >= SELL_ALL_SHARE * held[symbol]:
                    quantity = held[symbol]
                row.quantity = round(quantity, 6)
                row.price = round(row.amount_eur / row.quantity, 4)
            else:
                row.quantity = row.price = None
        if row.quantity:
            held[symbol] += row.quantity if row.kind == KIND_BUY else -row.quantity


def price_window(statement: Statement) -> Tuple[Optional[date], Optional[date]]:
    days = [r.day for r in statement.rows if r.quantity_estimated]
    if not days:
        return None, None
    return min(days) - timedelta(days=10), max(days) + timedelta(days=6)
