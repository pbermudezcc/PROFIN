"""
app/investments/tax_report.py
=============================
FAMILYFIN — Yearly tax summary of a user's investments (Spanish IRPF).

Orientative, not tax advice: it lays out what the return asks for so it
can be checked against the AEAT draft (datos fiscales), which brokers
outside Spain often do not report.

Sales ("ganancias y pérdidas patrimoniales")
    One line per FIFO lot a sale consumed: acquisition date and value,
    transmission date and value, result. Values are the EUR amounts of the
    operations (fees included in the purchase, deducted from the sale).

Anti-repurchase rule (art. 33.5.f LIRPF)
    A loss cannot be offset in that year while homogeneous securities
    bought within the window around the sale are still held. The window is
    two months for listed shares/ETFs and one year for (unlisted) funds.
    Repurchased shares = shares bought in the window before the sale that
    remain after it + shares bought in the window after it. They block the
    sale's losing lots in FIFO order, up to their quantity. The report only
    flags them: the deferred loss is recovered when those shares are sold.

Dividends ("rendimientos del capital mobiliario")
    Gross = net received + withholding at source; the withholding is the
    input for the international double-taxation deduction.
"""
from __future__ import annotations

import calendar
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Dict, List, Optional

from sqlalchemy.orm import Session

from app.db.models import AssetType, InvestmentKind, InvestmentTransaction, Security
from app.investments.portfolio import build_positions


def add_months(day: date, months: int) -> date:
    month_index = day.month - 1 + months
    year = day.year + month_index // 12
    month = month_index % 12 + 1
    return date(year, month, min(day.day, calendar.monthrange(year, month)[1]))


def window_months(security: Security) -> int:
    return 12 if AssetType(security.asset_type) == AssetType.fund else 2


@dataclass
class SaleLine:
    security:        Security
    broker:          str
    sell_date:       date
    acquired:        date
    quantity:        float
    proceeds_eur:    float
    cost_eur:        float
    blocked_eur:     float = 0.0            # part of the loss deferred (<= 0)
    repurchases:     List[date] = field(default_factory=list)

    @property
    def gain_eur(self) -> float:
        return self.proceeds_eur - self.cost_eur


@dataclass
class DividendLine:
    security:     Security
    broker:       str
    pay_date:     date
    net_eur:      float
    withheld_eur: float

    @property
    def gross_eur(self) -> float:
        return self.net_eur + self.withheld_eur


@dataclass
class TaxReport:
    year:      int
    sales:     List[SaleLine] = field(default_factory=list)
    dividends: List[DividendLine] = field(default_factory=list)

    @property
    def proceeds_eur(self) -> float:
        return sum(s.proceeds_eur for s in self.sales)

    @property
    def cost_eur(self) -> float:
        return sum(s.cost_eur for s in self.sales)

    @property
    def gains_eur(self) -> float:
        return sum(s.gain_eur for s in self.sales if s.gain_eur > 0)

    @property
    def losses_eur(self) -> float:
        return sum(s.gain_eur for s in self.sales if s.gain_eur < 0)

    @property
    def blocked_losses_eur(self) -> float:
        return sum(s.blocked_eur for s in self.sales)

    @property
    def net_eur(self) -> float:
        return self.gains_eur + self.losses_eur

    @property
    def computable_net_eur(self) -> float:
        """Net result once the losses blocked by the two-month rule are left out."""
        return self.net_eur - self.blocked_losses_eur

    @property
    def dividends_gross_eur(self) -> float:
        return sum(d.gross_eur for d in self.dividends)

    @property
    def dividends_withheld_eur(self) -> float:
        return sum(d.withheld_eur for d in self.dividends)

    @property
    def dividends_net_eur(self) -> float:
        return sum(d.net_eur for d in self.dividends)


def report_years(session: Session, user_id: int) -> List[int]:
    """Years with sales or dividends, newest first."""
    rows = (session.query(InvestmentTransaction.trade_date)
            .filter(InvestmentTransaction.user_id == user_id,
                    InvestmentTransaction.kind.in_([InvestmentKind.sell, InvestmentKind.dividend]))
            .all())
    return sorted({d.year for (d,) in rows}, reverse=True)


def tax_report(session: Session, user_id: int, year: int) -> TaxReport:
    ledger = (session.query(InvestmentTransaction)
              .filter(InvestmentTransaction.user_id == user_id).all())
    securities: Dict[int, Security] = {
        s.id: s for s in session.query(Security)
        .filter(Security.id.in_({t.security_id for t in ledger}))}
    by_id = {t.id: t for t in ledger}
    report = TaxReport(year=year)

    for security_id, position in build_positions(ledger).items():
        security = securities[security_id]
        by_sale: Dict[int, List] = {}
        for realized in position.realized:
            if realized.sell_date.year == year:
                by_sale.setdefault(realized.sell_transaction_id, []).append(realized)
        for sale_id, lots in by_sale.items():
            sale = by_id.get(sale_id)
            lines = [SaleLine(security=security, broker=sale.broker if sale else "",
                              sell_date=r.sell_date, acquired=r.acquired, quantity=r.quantity,
                              proceeds_eur=r.proceeds_eur, cost_eur=r.cost_eur) for r in lots]
            if sale is not None and any(line.gain_eur < 0 for line in lines):
                _apply_repurchase_rule(lines, sale, ledger, security)
            report.sales.extend(lines)

    for txn in ledger:
        if InvestmentKind(txn.kind) == InvestmentKind.dividend and txn.trade_date.year == year:
            report.dividends.append(DividendLine(
                security=securities[txn.security_id], broker=txn.broker, pay_date=txn.trade_date,
                net_eur=float(txn.amount_eur), withheld_eur=float(txn.tax_withheld_eur or 0.0)))

    report.sales.sort(key=lambda s: (s.sell_date, s.security.name, s.acquired))
    report.dividends.sort(key=lambda d: (d.pay_date, d.security.name))
    return report


def _apply_repurchase_rule(lines: List[SaleLine], sale, ledger, security: Security) -> None:
    months = window_months(security)
    start = add_months(sale.trade_date, -months)
    end = add_months(sale.trade_date, months)
    same = [t for t in ledger if t.security_id == security.id]

    # Bought in the window before the sale and still held right after it:
    # replay everything up to and including the sale (FIFO same-day order).
    through_sale = [t for t in same
                    if (t.trade_date, _order(t)) <= (sale.trade_date, _order(sale))]
    position = build_positions(through_sale).get(security.id)
    remaining = [lot for lot in (position.lots if position else []) if lot.acquired >= start]
    before_qty = sum(lot.quantity for lot in remaining)
    before_dates = [lot.acquired for lot in remaining]

    # Bought in the window after the sale.
    after = [t for t in same if sale.trade_date < t.trade_date <= end
             and InvestmentKind(t.kind) in (InvestmentKind.buy, InvestmentKind.grant)]
    after_qty = sum(float(t.quantity or 0.0) for t in after)
    repurchase_dates = sorted(set(before_dates + [t.trade_date for t in after]))

    available = before_qty + after_qty
    for line in lines:
        if line.gain_eur >= 0 or available <= 1e-9:
            continue
        blocked_qty = min(line.quantity, available)
        line.blocked_eur = line.gain_eur * blocked_qty / line.quantity
        line.repurchases = repurchase_dates
        available -= blocked_qty


def _order(txn) -> int:
    # Same-day replay order used by the FIFO engine (buys before sells).
    return {InvestmentKind.split: 0, InvestmentKind.buy: 1, InvestmentKind.grant: 1,
            InvestmentKind.sell: 2, InvestmentKind.dividend: 3}[InvestmentKind(txn.kind)] * 10**9 + (txn.id or 0)


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------
def export_xlsx(report: TaxReport, path: Path, owner: Optional[str] = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    money = '#,##0.00 "€";[Red]-#,##0.00 "€"'
    header_fill = PatternFill("solid", fgColor="1E3A8A")
    header_font = Font(bold=True, color="FFFFFF")
    wb = Workbook()

    summary = wb.active
    summary.title = "Resumen"
    title = f"Informe fiscal de inversiones {report.year}" + (f" — {owner}" if owner else "")
    summary.append([title])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append(["Orientativo: compruébalo con el borrador / datos fiscales de la AEAT."])
    summary.append([])
    rows = [
        ("Ganancias y pérdidas patrimoniales", None),
        ("Valor de transmisión", report.proceeds_eur),
        ("Valor de adquisición", report.cost_eur),
        ("Ganancias", report.gains_eur),
        ("Pérdidas", report.losses_eur),
        ("Resultado neto", report.net_eur),
        ("Pérdidas no computables este año (regla de los 2 meses / 1 año)", report.blocked_losses_eur),
        ("Resultado computable", report.computable_net_eur),
        (None, None),
        ("Dividendos (rendimientos del capital mobiliario)", None),
        ("Dividendos íntegros (brutos)", report.dividends_gross_eur),
        ("Retenciones en origen (deducción doble imposición internacional)", report.dividends_withheld_eur),
        ("Dividendos netos cobrados", report.dividends_net_eur),
    ]
    for label, value in rows:
        summary.append([label, round(value, 2) if value is not None else None])
        row = summary.max_row
        if value is None and label:
            summary.cell(row, 1).font = Font(bold=True)
        else:
            summary.cell(row, 2).number_format = money
    summary.column_dimensions["A"].width = 66
    summary.column_dimensions["B"].width = 18

    def table(sheet, headers, data, money_cols, date_cols, qty_cols=()):
        sheet.append(headers)
        for col in range(1, len(headers) + 1):
            cell = sheet.cell(1, col)
            cell.fill, cell.font = header_fill, header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for values in data:
            sheet.append(values)
            row = sheet.max_row
            for col in money_cols:
                sheet.cell(row, col).number_format = money
            for col in date_cols:
                sheet.cell(row, col).number_format = "DD/MM/YYYY"
            for col in qty_cols:
                sheet.cell(row, col).number_format = "#,##0.######"
        for col, header in enumerate(headers, start=1):
            longest = max([len(str(header))] + [len(str(v[col - 1])) for v in data if v[col - 1] is not None])
            sheet.column_dimensions[get_column_letter(col)].width = min(max(12, longest + 2), 48)
        sheet.freeze_panes = "A2"

    sales = wb.create_sheet("Ganancias y pérdidas")
    table(sales,
          ["Valor", "Símbolo", "ISIN", "Bróker", "Fecha de transmisión", "Fecha de adquisición",
           "Acciones", "Valor de transmisión", "Valor de adquisición", "Resultado",
           "Pérdida no computable", "Recompras en el plazo"],
          [[s.security.name, s.security.symbol, s.security.isin or "", s.broker, s.sell_date,
            s.acquired, s.quantity, round(s.proceeds_eur, 2), round(s.cost_eur, 2),
            round(s.gain_eur, 2), round(s.blocked_eur, 2) if s.blocked_eur else None,
            ", ".join(d.strftime("%d/%m/%Y") for d in s.repurchases)]
           for s in report.sales],
          money_cols=(8, 9, 10, 11), date_cols=(5, 6), qty_cols=(7,))

    dividends = wb.create_sheet("Dividendos")
    table(dividends,
          ["Valor", "Símbolo", "ISIN", "Bróker", "Fecha de cobro", "Íntegro", "Retención en origen",
           "Neto cobrado"],
          [[d.security.name, d.security.symbol, d.security.isin or "", d.broker, d.pay_date,
            round(d.gross_eur, 2), round(d.withheld_eur, 2), round(d.net_eur, 2)]
           for d in report.dividends],
          money_cols=(6, 7, 8), date_cols=(5,))

    path = Path(path)
    wb.save(path)
    return path
