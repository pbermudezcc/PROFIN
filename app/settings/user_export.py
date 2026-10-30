"""
app/settings/user_export.py
===========================
FAMILYFIN — «Descargar mis datos»: one member's own data as a zip (Qt-free,
shared by the web «Mi cuenta» and the PC account menu).

  export_user_data(session, user_id, when=None) -> UserExport

The zip holds:
  * FAMILYFIN-<usuario>-<fecha>.xlsx — Movimientos, Cuentas, Categorías,
    Reglas, Operaciones, Efectivo and Importaciones sheets;
  * the categories and personal rules as a .profin-cats.json library, which
    «Importar» (Categorías) reads back;
  * LEEME.txt.

Only rows of this user: never the password hash, the family's global rules,
the audit log or another member's data. It is an export to keep or to open
in Excel, not a backup the app can restore: the full family backup
(SettingsManager.backup_all_data) stays on the PC and is never offered on
the web.
"""
from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Dict, List, Optional, Sequence

from sqlalchemy.orm import Session, joinedload

from app.analytics.transaction_search import SEARCH_TYPE_LABELS
from app.categorization import library_io
from app.db.models import (
    CATEGORY_SOURCE_MANUAL,
    CATEGORY_SOURCE_RULE,
    TRANSACTIONS_NEWEST_FIRST,
    Account,
    AutoCategorizationRule,
    BrokerCash,
    Category,
    CategoryLevel,
    ImportBatch,
    InvestmentTransaction,
    Transaction,
    User,
)
from app.investments.ledger import KIND_LABELS

MATCH_LABELS = {"contains": "Contiene", "exact": "Exacta", "regex": "Expresión regular"}
SOURCE_LABELS = {CATEGORY_SOURCE_MANUAL: "Manual", CATEGORY_SOURCE_RULE: "Regla automática"}
STATUS_LABELS = {"pending": "Pendiente", "completed": "Completada", "failed": "Fallida", "partial": "Parcial"}


@dataclass
class UserExport:
    data: bytes
    filename: str
    counts: Dict[str, int] = field(default_factory=dict)

    def describe(self) -> str:
        """Short Spanish summary for the audit log and the PC message."""
        c = self.counts
        return (f"{c.get('movimientos', 0)} movimientos, {c.get('operaciones', 0)} operaciones, "
                f"{c.get('categorias', 0)} categorías, {c.get('reglas', 0)} reglas")


def safe_name(username: str) -> str:
    """ASCII, file-name-safe form of a user name («Lucía» → «lucia»)."""
    folded = unicodedata.normalize("NFKD", username).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", folded.lower()).strip("-") or "usuario"


def export_filename(username: str, when: Optional[datetime] = None) -> str:
    return f"familyfin-datos-{safe_name(username)}-{(when or datetime.now()):%Y%m%d}.zip"


def _naive(value):
    """openpyxl refuses timezone-aware datetimes."""
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    return value


def _label(enum_value, labels: Dict) -> str:
    if enum_value is None:
        return ""
    return labels.get(enum_value, labels.get(getattr(enum_value, "value", None), str(enum_value)))


def _workbook(sheets: Sequence[tuple]) -> bytes:
    """sheets: (title, headers, rows, money columns, date columns)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    money = '#,##0.00 "€";[Red]-#,##0.00 "€"'
    header_fill = PatternFill("solid", fgColor="1E3A8A")
    header_font = Font(bold=True, color="FFFFFF")
    wb = Workbook()
    wb.remove(wb.active)
    for title, headers, rows, money_cols, date_cols in sheets:
        sheet = wb.create_sheet(title)
        sheet.append(headers)
        for col in range(1, len(headers) + 1):
            cell = sheet.cell(1, col)
            cell.fill, cell.font = header_fill, header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        widths = [len(h) for h in headers]
        for values in rows:
            sheet.append([_naive(v) for v in values])
            row = sheet.max_row
            for col in money_cols:
                sheet.cell(row, col).number_format = money
            for col in date_cols:
                sheet.cell(row, col).number_format = "DD/MM/YYYY"
            for i, v in enumerate(values):
                if v is not None:
                    widths[i] = max(widths[i], len(str(v)))
        for col, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(col)].width = min(max(12, width + 2), 60)
        sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _readme(user: User, when: datetime, counts: Dict[str, int], xlsx_name: str, library_name: str) -> str:
    return "\n".join([
        f"FAMILYFIN — datos de «{user.username}»",
        f"Exportado el {when:%d/%m/%Y a las %H:%M}.",
        "",
        "Contenido:",
        f"  {xlsx_name}",
        f"    Movimientos ({counts['movimientos']}), Cuentas ({counts['cuentas']}), "
        f"Categorías ({counts['categorias']}), Reglas ({counts['reglas']}),",
        f"    Operaciones de inversión ({counts['operaciones']}), Efectivo en el bróker "
        f"({counts['efectivo']}) e Importaciones ({counts['importaciones']}).",
        f"  {library_name}",
        "    Tus categorías y reglas personales. Se pueden volver a cargar en",
        "    FAMILYFIN desde Categorías › Importar (en el PC o en la web).",
        "",
        "Solo contiene tus datos: ni los de otros miembros de la familia, ni las",
        "reglas globales, ni tu contraseña.",
        "",
        "Esto es una exportación para guardar o abrir con Excel, no una copia de",
        "seguridad que la aplicación pueda restaurar. Las copias de seguridad de",
        "toda la familia se hacen en el PC (Configuración › Copias de seguridad).",
        "",
        "Guárdalo en un sitio seguro: incluye tus movimientos bancarios.",
        "",
    ])


def export_user_data(session: Session, user_id: int, when: Optional[datetime] = None) -> UserExport:
    """Build the zip of one user's own data (nothing is written to disk)."""
    user = session.get(User, user_id)
    if user is None:
        raise ValueError("La cuenta ya no existe.")
    when = when or datetime.now()

    accounts = (session.query(Account).filter(Account.user_id == user_id)
                .order_by(Account.id).all())
    account_names = {a.id: a.account_name for a in accounts}
    categories = session.query(Category).filter(Category.user_id == user_id).all()
    by_id = {c.id: c for c in categories}

    def category_pair(category_id: Optional[int]) -> tuple:
        cat = by_id.get(category_id)
        if cat is None:
            return "", ""
        parent = by_id.get(cat.parent_id)
        if parent is None:
            return cat.name, ""
        return parent.name, cat.name

    movements = (
        session.query(Transaction)
        .filter(Transaction.user_id == user_id)
        .order_by(*TRANSACTIONS_NEWEST_FIRST)
        .all()
    )
    movement_rows = [
        [t.transaction_date, t.value_date, account_names.get(t.account_id, ""),
         t.issuer_name, t.concept, t.concepto_comun, t.concepto_propio, t.raw_details,
         t.amount, t.running_balance, _label(t.type, SEARCH_TYPE_LABELS),
         *category_pair(t.category_id), SOURCE_LABELS.get(t.category_source, t.category_source or ""),
         t.source_file]
        for t in movements
    ]

    account_rows = [[a.account_name, a.bank, a.iban_last4, a.created_at] for a in accounts]

    category_rows: List[list] = []
    for root in sorted((c for c in categories if c.parent_id is None), key=lambda c: c.name.lower()):
        category_rows.append([root.name, "", _label(root.transaction_type, SEARCH_TYPE_LABELS)])
        for child in sorted((c for c in categories if c.parent_id == root.id), key=lambda c: c.name.lower()):
            category_rows.append([root.name, child.name, ""])

    rules = (
        session.query(AutoCategorizationRule)
        .filter(AutoCategorizationRule.user_id == user_id)
        .order_by(AutoCategorizationRule.priority, AutoCategorizationRule.id)
        .all()
    )
    rule_rows = [
        [r.priority, r.match_pattern, _label(r.match_type, MATCH_LABELS), *category_pair(r.category_id), r.created_at]
        for r in rules
    ]

    operations = (
        session.query(InvestmentTransaction)
        .options(joinedload(InvestmentTransaction.security))
        .filter(InvestmentTransaction.user_id == user_id)
        .order_by(InvestmentTransaction.trade_date.desc(), InvestmentTransaction.id.desc())
        .all()
    )
    operation_rows = [
        [o.trade_date, o.broker, _label(o.kind, KIND_LABELS), o.security.name, o.security.symbol,
         o.security.isin, o.quantity, o.price, o.currency, o.amount_eur, o.fees_eur,
         o.tax_withheld_eur, o.cash_paid_eur, o.split_ratio, o.notes, o.broker_ref]
        for o in operations
    ]

    cash = (session.query(BrokerCash).filter(BrokerCash.user_id == user_id)
            .order_by(BrokerCash.broker, BrokerCash.currency).all())
    cash_rows = [[c.broker, c.currency, c.amount, c.as_of, c.source] for c in cash]

    batches = (session.query(ImportBatch).filter(ImportBatch.user_id == user_id)
               .order_by(ImportBatch.import_timestamp.desc()).all())
    batch_rows = [[b.filename, b.import_timestamp, b.rows_imported, b.rows_skipped,
                   _label(b.status, STATUS_LABELS)] for b in batches]

    counts = {
        "movimientos": len(movement_rows), "cuentas": len(account_rows),
        "categorias": sum(1 for c in categories if c.level == CategoryLevel.category),
        "reglas": len(rule_rows), "operaciones": len(operation_rows),
        "efectivo": len(cash_rows), "importaciones": len(batch_rows),
    }

    xlsx = _workbook([
        ("Movimientos",
         ["Fecha", "Fecha valor", "Cuenta", "Emisor", "Concepto", "Concepto común", "Concepto propio",
          "Detalles", "Importe", "Saldo", "Tipo", "Tipo de categoría", "Categoría",
          "Categorizado por", "Extracto"],
         movement_rows, (9, 10), (1, 2)),
        ("Cuentas", ["Cuenta", "Banco", "IBAN (últimos 4)", "Creada"], account_rows, (), ()),
        ("Categorías", ["Tipo", "Categoría", "Tipo de movimiento"], category_rows, (), ()),
        ("Reglas", ["Prioridad", "Patrón", "Coincidencia", "Tipo", "Categoría", "Creada"], rule_rows, (), ()),
        ("Operaciones",
         ["Fecha", "Bróker", "Operación", "Valor", "Símbolo", "ISIN", "Cantidad", "Precio", "Divisa",
          "Importe (EUR)", "Comisiones (EUR)", "Retención (EUR)", "Efectivo pagado (EUR)",
          "Ratio del split", "Notas", "Referencia del bróker"],
         operation_rows, (10, 11, 12, 13), (1,)),
        ("Efectivo", ["Bróker", "Divisa", "Importe", "A fecha", "Origen"], cash_rows, (), (4,)),
        ("Importaciones", ["Extracto", "Importado el", "Filas importadas", "Filas omitidas", "Estado"],
         batch_rows, (), ()),
    ])

    stem = f"FAMILYFIN-{safe_name(user.username)}-{when:%Y%m%d}"
    xlsx_name = f"{stem}.xlsx"
    library_name = f"profin-categorias-{safe_name(user.username)}-{when:%Y%m%d}{library_io.FILE_SUFFIX}"
    library = library_io.dumps(library_io.export_library(session, user_id, include_global=False,
                                                         source_user=user.username))

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(xlsx_name, xlsx)
        archive.writestr(library_name, library)
        archive.writestr("LEEME.txt", _readme(user, when, counts, xlsx_name, library_name).encode("utf-8"))
    return UserExport(data=buffer.getvalue(), filename=export_filename(user.username, when), counts=counts)
