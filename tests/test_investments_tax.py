"""
tests/test_investments_tax.py
=============================
Pending dividends/splits detection, split-adjusted history, the yearly tax
report (FIFO lots, anti-repurchase rule, dividends) and its Excel export.
"""
from __future__ import annotations

from datetime import date
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.db.database import _build_engine
from app.db.models import (
    PRICE_SOURCE_MANUAL,
    AssetType,
    Base,
    FxRate,
    InvestmentKind,
    PriceHistory,
    Security,
    User,
)
from app.investments.events import pending_events
from app.investments.ledger import TransactionDraft, delete_transaction, save_transaction
from app.investments.market_data import CorporateActions
from app.investments.performance import portfolio_history
from app.investments.prices import set_manual_price
from app.investments.tax_report import add_months, export_xlsx, report_years, tax_report

B, S, D, SP = (InvestmentKind.buy, InvestmentKind.sell, InvestmentKind.dividend,
               InvestmentKind.split)


@pytest.fixture
def db(tmp_path):
    engine = _build_engine(f"sqlite:///{tmp_path / 'tax.db'}")
    Base.metadata.create_all(engine)
    s = Session(engine)
    user = User(username="pedro", password_hash="x")
    other = User(username="maria", password_hash="x")
    gev = Security(symbol="GEV", name="GE Vernova Inc.", currency="USD", isin="US36828A1016")
    enr = Security(symbol="ENR.DE", name="Siemens Energy AG", currency="EUR")
    fund = Security(symbol="0P00008XG7.F", name="Boreas", currency="EUR", asset_type=AssetType.fund)
    s.add_all([user, other, gev, enr, fund])
    s.flush()
    s.add(FxRate(currency="USD", rate_date=date(2026, 1, 2), rate=1.25))
    s.commit()
    yield SimpleNamespace(s=s, user=user, other=other, gev=gev, enr=enr, fund=fund,
                          path=tmp_path)
    s.close()
    engine.dispose()


def op(db, sec, kind, d, qty=None, eur=0.0, user=None, **kw):
    return save_transaction(db.s, (user or db.user).id, TransactionDraft(
        security_id=sec.id, kind=kind, trade_date=d, broker=kw.pop("broker", "DEGIRO"),
        quantity=qty, amount_eur=eur, **kw))


# ---------------------------------------------------------------------------
# Pending corporate actions
# ---------------------------------------------------------------------------

def test_unregistered_dividends_and_splits_are_proposed(db):
    op(db, db.gev, B, date(2025, 3, 3), 4, 1000.0)
    actions = {"GEV": CorporateActions(
        dividends=((date(2024, 12, 20), 0.25), (date(2026, 1, 5), 0.5), (date(2026, 3, 17), 0.5)),
        splits=((date(2026, 2, 1), 2.0),))}
    op(db, db.gev, D, date(2026, 3, 30), eur=1.36, tax_withheld_eur=0.24)   # pays the March one
    events = pending_events(db.s, db.user.id, actions, today=date(2026, 9, 27))
    assert [(e.kind, e.ex_date) for e in events] == [
        (SP, date(2026, 2, 1)), (D, date(2026, 1, 5))]
    split, dividend = events
    assert split.ratio == 2.0 and split.shares == 4
    assert dividend.shares == 4 and dividend.gross_eur == pytest.approx(4 * 0.5 / 1.25)
    assert dividend.key == "dividend|GEV|2026-01-05"
    # Dismissed ones disappear; before the first purchase nothing is proposed.
    assert pending_events(db.s, db.user.id, actions, dismissed={dividend.key, split.key},
                          today=date(2026, 9, 27)) == []


def test_no_dividend_proposed_when_nothing_was_held_or_for_other_users(db):
    op(db, db.enr, B, date(2025, 1, 2), 10, 1000.0)
    op(db, db.enr, S, date(2025, 6, 1), 10, 1200.0)
    actions = {"ENR.DE": CorporateActions(dividends=((date(2026, 2, 27), 0.7),))}
    assert pending_events(db.s, db.user.id, actions) == []
    assert pending_events(db.s, db.other.id, actions) == []


# ---------------------------------------------------------------------------
# Splits: history and price cache
# ---------------------------------------------------------------------------

def test_history_uses_split_adjusted_closes_correctly(db):
    op(db, db.enr, B, date(2026, 1, 5), 10, 1000.0)          # paid 100 €/share pre-split
    op(db, db.enr, SP, date(2026, 1, 7), split_ratio=2.0)      # resets the price cache
    # What the provider returns after a 2-for-1 split on 2026-01-07: earlier closes halved.
    for day, close in ((date(2026, 1, 5), 50.0), (date(2026, 1, 6), 55.0), (date(2026, 1, 7), 60.0)):
        db.s.add(PriceHistory(security_id=db.enr.id, price_date=day, close=close))
    db.s.commit()
    values = [round(p.value_eur, 2) for p in portfolio_history(db.s, db.user.id,
                                                               end=date(2026, 1, 7)).points]
    assert values == [1000.0, 1100.0, 1200.0]


def test_saving_or_deleting_a_split_resets_downloaded_prices_only(db):
    db.s.add(PriceHistory(security_id=db.enr.id, price_date=date(2026, 1, 5), close=50.0))
    db.s.commit()
    set_manual_price(db.s, db.enr.id, date(2026, 1, 6), 51.0)
    op(db, db.enr, B, date(2026, 1, 5), 10, 1000.0)
    split = op(db, db.enr, SP, date(2026, 1, 7), split_ratio=2.0)
    rows = db.s.query(PriceHistory).filter_by(security_id=db.enr.id).all()
    assert [(r.price_date, r.source) for r in rows] == [(date(2026, 1, 6), PRICE_SOURCE_MANUAL)]
    db.s.add(PriceHistory(security_id=db.enr.id, price_date=date(2026, 1, 8), close=60.0))
    db.s.commit()
    delete_transaction(db.s, db.user.id, split.id)
    assert db.s.query(PriceHistory).filter_by(security_id=db.enr.id, source="yahoo").count() == 0


# ---------------------------------------------------------------------------
# Tax report
# ---------------------------------------------------------------------------

def test_add_months_clamps_end_of_month():
    assert add_months(date(2026, 3, 31), -1) == date(2026, 2, 28)
    assert add_months(date(2026, 11, 15), 2) == date(2027, 1, 15)


def test_sales_are_reported_per_fifo_lot_with_totals(db):
    op(db, db.enr, B, date(2024, 1, 10), 10, 1000.0)
    op(db, db.enr, B, date(2024, 6, 10), 10, 1500.0)
    op(db, db.enr, S, date(2025, 3, 1), 15, 2400.0)
    report = tax_report(db.s, db.user.id, 2025)
    assert [(s.acquired, s.quantity) for s in report.sales] == [
        (date(2024, 1, 10), 10), (date(2024, 6, 10), 5)]
    assert report.proceeds_eur == pytest.approx(2400.0)
    assert report.cost_eur == pytest.approx(1750.0)
    assert report.gains_eur == pytest.approx(650.0) and report.losses_eur == 0
    assert report.blocked_losses_eur == 0 and report.computable_net_eur == pytest.approx(650.0)
    assert tax_report(db.s, db.user.id, 2024).sales == []


def test_loss_with_repurchase_within_two_months_is_blocked(db):
    op(db, db.enr, B, date(2025, 1, 10), 10, 1500.0)
    op(db, db.enr, S, date(2025, 5, 2), 10, 1000.0)          # -500
    op(db, db.enr, B, date(2025, 6, 20), 4, 420.0)           # 49 days later
    report = tax_report(db.s, db.user.id, 2025)
    line = report.sales[0]
    assert line.gain_eur == pytest.approx(-500.0)
    assert line.blocked_eur == pytest.approx(-200.0)          # 4 of 10 shares repurchased
    assert line.repurchases == [date(2025, 6, 20)]
    assert report.net_eur == pytest.approx(-500.0)
    assert report.computable_net_eur == pytest.approx(-300.0)


def test_shares_bought_just_before_and_still_held_block_the_loss(db):
    op(db, db.enr, B, date(2024, 1, 10), 10, 1500.0)
    op(db, db.enr, B, date(2025, 4, 1), 10, 1000.0)          # within 2 months before the sale
    op(db, db.enr, S, date(2025, 5, 2), 10, 1000.0)          # FIFO sells the 2024 lot: -500
    line = tax_report(db.s, db.user.id, 2025).sales[0]
    assert line.blocked_eur == pytest.approx(-500.0)
    assert line.repurchases == [date(2025, 4, 1)]


def test_repurchase_outside_the_window_or_a_gain_is_not_flagged(db):
    op(db, db.enr, B, date(2025, 1, 10), 10, 1500.0)
    op(db, db.enr, S, date(2025, 5, 2), 10, 1000.0)
    op(db, db.enr, B, date(2025, 7, 10), 4, 420.0)           # > 2 months later
    assert tax_report(db.s, db.user.id, 2025).sales[0].blocked_eur == 0


def test_funds_use_a_one_year_window(db):
    op(db, db.fund, B, date(2025, 1, 10), 10, 1500.0, broker="Tressis")
    op(db, db.fund, S, date(2025, 5, 2), 10, 1000.0, broker="Tressis")
    op(db, db.fund, B, date(2025, 12, 1), 10, 900.0, broker="Tressis")
    assert tax_report(db.s, db.user.id, 2025).sales[0].blocked_eur == pytest.approx(-500.0)


def test_dividends_gross_withholding_and_years(db):
    op(db, db.gev, B, date(2025, 3, 3), 4, 1000.0)
    op(db, db.gev, D, date(2025, 4, 30), eur=8.5, tax_withheld_eur=1.5)
    op(db, db.gev, D, date(2026, 1, 20), eur=1.7, tax_withheld_eur=0.3)
    report = tax_report(db.s, db.user.id, 2025)
    assert (report.dividends_gross_eur, report.dividends_withheld_eur, report.dividends_net_eur) == (
        pytest.approx(10.0), pytest.approx(1.5), pytest.approx(8.5))
    assert report_years(db.s, db.user.id) == [2026, 2025]
    assert report_years(db.s, db.other.id) == []


def test_excel_export_has_three_sheets_with_values(db):
    op(db, db.enr, B, date(2025, 1, 10), 10, 1500.0)
    op(db, db.enr, S, date(2025, 5, 2), 10, 1000.0)
    op(db, db.enr, B, date(2025, 6, 20), 4, 420.0)
    op(db, db.gev, B, date(2025, 3, 3), 4, 1000.0)
    op(db, db.gev, D, date(2025, 4, 30), eur=8.5, tax_withheld_eur=1.5)
    path = export_xlsx(tax_report(db.s, db.user.id, 2025), db.path / "informe.xlsx", owner="pedro")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Resumen", "Ganancias y pérdidas", "Dividendos"]
    summary = {row[0]: row[1] for row in wb["Resumen"].iter_rows(values_only=True) if row and row[0]}
    assert summary["Resultado neto"] == pytest.approx(-500.0)
    assert summary["Resultado computable"] == pytest.approx(-300.0)
    assert summary["Dividendos íntegros (brutos)"] == pytest.approx(10.0)
    sale = list(wb["Ganancias y pérdidas"].iter_rows(min_row=2, values_only=True))[0]
    assert sale[0] == "Siemens Energy AG" and sale[9] == pytest.approx(-500.0)
    assert sale[10] == pytest.approx(-200.0) and sale[11] == "20/06/2025"
    dividend = list(wb["Dividendos"].iter_rows(min_row=2, values_only=True))[0]
    assert dividend[2] == "US36828A1016" and dividend[5] == pytest.approx(10.0)
