"""
tests/test_user_export.py
=========================
«Descargar mis datos»: the zip holds only this member's rows, never the
password hash, the family's global rules or another member's data.
"""
from __future__ import annotations

import io
import json
import zipfile
from datetime import date, datetime

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.auth.auth_service import create_user
from app.categorization.category_manager import seed_default_categories, seed_default_rules
from app.db.models import (
    Account,
    AssetType,
    AutoCategorizationRule,
    Base,
    BrokerCash,
    Category,
    InvestmentKind,
    InvestmentTransaction,
    MatchType,
    Security,
    Transaction,
    TransactionType,
    User,
)
from app.settings.user_export import export_filename, export_user_data

WHEN = datetime(2026, 10, 30, 12, 0)


@pytest.fixture
def session():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    sess = sessionmaker(bind=engine)()
    yield sess
    sess.close()
    engine.dispose()


def _member(session, name, issuer):
    user = create_user(session, name, "clave-secreta-1")
    seed_default_categories(session, user.id)
    seed_default_rules(session, user.id)
    account = Account(user_id=user.id, account_name=f"Cuenta de {name}", iban_last4="1234")
    session.add(account)
    session.flush()
    luz = session.query(Category).filter_by(user_id=user.id, name="Luz").one()
    session.add(Transaction(user_id=user.id, account_id=account.id, transaction_date=date(2026, 9, 1),
                            amount=-45.3, type=TransactionType.expense, issuer_name=issuer,
                            category_id=luz.id, category_source="rule"))
    session.commit()
    return user


@pytest.fixture
def family(session):
    ana = _member(session, "Lucía", "IBERDROLA ANA")
    beto = _member(session, "beto", "SECRETO DE BETO")
    vwce = Security(symbol="VWCE.DE", name="Vanguard FTSE All-World", asset_type=AssetType.etf,
                    currency="EUR", isin="IE00BK5BQT80")
    session.add(vwce)
    session.flush()
    session.add(InvestmentTransaction(user_id=ana.id, security_id=vwce.id, broker="DEGIRO",
                                      kind=InvestmentKind.buy, trade_date=date(2026, 9, 2),
                                      quantity=2, price=120.0, currency="EUR", amount_eur=240.0,
                                      fees_eur=1.0, tax_withheld_eur=0.0))
    session.add(BrokerCash(user_id=ana.id, broker="DEGIRO", currency="EUR", amount=1002.42,
                           as_of=date(2026, 10, 27), source="DEGIRO"))
    luz = session.query(Category).filter_by(user_id=ana.id, name="Luz").one()
    session.add(AutoCategorizationRule(user_id=None, category_id=luz.id,
                                       match_pattern="REGLA FAMILIAR", match_type=MatchType.contains))
    session.commit()
    return ana, beto


def _open(export):
    archive = zipfile.ZipFile(io.BytesIO(export.data))
    names = archive.namelist()
    xlsx = next(n for n in names if n.endswith(".xlsx"))
    wb = load_workbook(io.BytesIO(archive.read(xlsx)))
    return archive, names, wb


def test_zip_holds_only_this_members_data(session, family):
    ana, beto = family

    export = export_user_data(session, ana.id, when=WHEN)
    archive, names, wb = _open(export)

    assert export.filename == "familyfin-datos-lucia-20261030.zip"
    assert sorted(names) == sorted(["FAMILYFIN-lucia-20261030.xlsx", "LEEME.txt",
                                    "profin-categorias-lucia-20261030.profin-cats.json"])
    assert wb.sheetnames == ["Movimientos", "Cuentas", "Categorías", "Reglas", "Operaciones",
                             "Efectivo", "Importaciones"]

    movements = list(wb["Movimientos"].iter_rows(min_row=2, values_only=True))
    assert len(movements) == 1
    row = movements[0]
    assert row[0].date() == date(2026, 9, 1) and row[3] == "IBERDROLA ANA" and row[8] == -45.3
    assert (row[10], row[11], row[12], row[13]) == ("Gasto", "GASTO", "Luz", "Regla automática")

    ops = list(wb["Operaciones"].iter_rows(min_row=2, values_only=True))
    assert [(o[2], o[4], o[9]) for o in ops] == [("Compra", "VWCE.DE", 240.0)]
    assert list(wb["Efectivo"].iter_rows(min_row=2, values_only=True))[0][2] == 1002.42

    everything = b"".join(archive.read(n) for n in names if not n.endswith(".xlsx"))
    cells = " ".join(str(c) for ws in wb for r in ws.iter_rows(values_only=True) for c in r if c)
    for secret in ("SECRETO DE BETO", "beto", "REGLA FAMILIAR", session.get(User, ana.id).password_hash):
        assert secret not in cells
        assert secret.encode() not in everything

    library = json.loads(archive.read("profin-categorias-lucia-20261030.profin-cats.json"))
    assert all(rule["scope"] == "personal" for rule in library["rules"])
    assert "Lucía" in archive.read("LEEME.txt").decode("utf-8")
    assert export.counts["movimientos"] == 1 and export.counts["operaciones"] == 1


def test_empty_account_exports_cleanly(session):
    user = create_user(session, "nuevo", "clave-secreta-1")

    export = export_user_data(session, user.id, when=WHEN)
    _, _, wb = _open(export)

    assert all(ws.max_row == 1 for ws in wb)          # headers only
    assert export.describe().startswith("0 movimientos")


def test_file_name_is_ascii():
    assert export_filename("José Mª Ñ", WHEN) == "familyfin-datos-jose-ma-n-20261030.zip"
