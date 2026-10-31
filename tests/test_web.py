"""
tests/test_web.py
=================
LAN web front-end: authentication, per-user isolation, CSRF/origin checks,
security headers, pages, light edits and the shared services behind them.
"""
from __future__ import annotations

import re
from datetime import date

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event
from sqlalchemy.orm import sessionmaker

from app.analytics.dashboard_summary import dashboard_summary
from app.auth.auth_service import create_user
from app.categorization.category_manager import assign_user_transaction_category
from app.db.models import (
    CATEGORY_SOURCE_MANUAL,
    Account,
    AssetType,
    AuditLog,
    AutoCategorizationRule,
    Base,
    Category,
    CategoryLevel,
    ImportBatch,
    InvestmentKind,
    InvestmentTransaction,
    MatchType,
    Security,
    Transaction,
    TransactionType,
)
from app.web.app import create_app
from app.web.deps import get_db
from app.web.security import SESSION_COOKIE, LoginLimiter, SessionStore

BASE = "https://testserver"
ORIGIN = {"Origin": BASE}


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture
def factory(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'web.db'}", connect_args={"check_same_thread": False})
    event.listen(engine, "connect", lambda c, _r: c.execute("PRAGMA foreign_keys = ON"))
    Base.metadata.create_all(engine)
    yield sessionmaker(bind=engine, autoflush=False)
    engine.dispose()


@pytest.fixture
def family(factory):
    db = factory()
    ana = create_user(db, "ana", "clave-ana")
    beto = create_user(db, "beto", "clave-beto")
    ids = {"ana": ana.id, "beto": beto.id}
    for user in (ana, beto):
        account = Account(user_id=user.id, account_name=f"Cuenta {user.username}")
        db.add(account)
        db.flush()
        gasto = Category(user_id=user.id, name="GASTO", level=CategoryLevel.type,
                         transaction_type=TransactionType.expense)
        db.add(gasto)
        db.flush()
        super_ = Category(user_id=user.id, parent_id=gasto.id, name="Supermercados",
                          level=CategoryLevel.category)
        db.add(super_)
        db.flush()
        ids[f"{user.username}_cat"] = super_.id
        txn = Transaction(user_id=user.id, account_id=account.id, transaction_date=date.today(),
                          amount=-42.5, running_balance=1000.0, type=TransactionType.expense,
                          issuer_name=f"MERCADO {user.username.upper()}")
        nomina = Transaction(user_id=user.id, account_id=account.id, transaction_date=date.today(),
                             amount=2000.0, type=TransactionType.income, issuer_name="NOMINA")
        db.add_all([txn, nomina])
        db.flush()
        ids[f"{user.username}_txn"] = txn.id
    security = Security(symbol="ACME", name="Acme Corp", asset_type=AssetType.stock, currency="EUR")
    db.add(security)
    db.flush()
    buy = InvestmentTransaction(user_id=beto.id, security_id=security.id, broker="DEGIRO",
                                kind=InvestmentKind.buy, trade_date=date(2025, 1, 10),
                                quantity=10, price=10, currency="EUR", amount_eur=100)
    db.add(buy)
    db.commit()
    ids["security"] = security.id
    ids["beto_op"] = buy.id
    db.close()
    return ids


class OfflineProvider:
    """Default market data in these tests: never the network, nothing found."""
    name = "offline"

    def search(self, text, limit=10):
        return []

    def daily_closes(self, symbols, start, end=None):
        return {}

    def security_info(self, symbol):
        from app.investments.market_data import MarketDataError
        raise MarketDataError("sin conexión en los tests")


@pytest.fixture
def app(factory, family, tmp_path):
    changes = []
    backups = []

    def fake_backup(user_name):
        # Never the real backup folder: a file in this test's tmp_path.
        path = tmp_path / f"profin_predelete_{user_name}.zip"
        path.write_bytes(b"copia")
        backups.append(user_name)
        return path

    application = create_app(sessions=SessionStore(), limiter=LoginLimiter(),
                             on_change=lambda uid, area: changes.append((uid, area)), lan_only=True,
                             market_provider=OfflineProvider(), backup_before_delete=fake_backup)
    application.state.backups = backups

    def _db():
        session = factory()
        try:
            yield session
        finally:
            session.close()

    application.dependency_overrides[get_db] = _db
    application.state.session_factory = factory          # background threads (price refresh)
    application.state.changes = changes
    yield application
    application.state.imports.clear()       # temp folders of unconfirmed uploads
    application.state.investment_imports.close()
    application.state.price_refresh.wait(timeout=10)     # never outlive the test database
    application.state.price_refresh.close()


def client_for(app) -> TestClient:
    # TestClient's peer address is "testclient"; present a LAN address.
    return TestClient(app, base_url=BASE, client=("192.168.1.50", 50000))


def login(client: TestClient, username: str, password: str):
    return client.post("/login", data={"username": username, "password": password},
                       headers=ORIGIN, follow_redirects=False)


def logged_in(app, username="ana", password="clave-ana") -> TestClient:
    client = client_for(app)
    response = login(client, username, password)
    assert response.status_code == 303, response.text
    return client


def csrf_of(client: TestClient) -> str:
    page = client.get("/").text
    return re.search(r'name="csrf-token" content="([^"]+)"', page).group(1)


# ---------------------------------------------------------------------------
# Authentication
# ---------------------------------------------------------------------------

def test_pages_require_login(app):
    client = client_for(app)
    response = client.get("/", follow_redirects=False)
    assert response.status_code == 303 and response.headers["location"] == "/login"
    assert client.get("/api/resumen").status_code == 401


def test_login_sets_hardened_session_cookie(app):
    client = client_for(app)
    response = login(client, "ana", "clave-ana")
    assert response.status_code == 303
    cookie = response.headers["set-cookie"]
    assert cookie.startswith(f"{SESSION_COOKIE}=")
    for flag in ("HttpOnly", "Secure", "SameSite=strict", "Path=/"):
        assert flag.lower() in cookie.lower()
    assert "domain" not in cookie.lower()
    assert "Hola, ana" in client.get("/").text


def test_wrong_password_is_generic_and_rate_limited(app):
    client = client_for(app)
    bad = login(client, "ana", "nope")
    assert bad.status_code == 401
    assert "Usuario o contraseña incorrectos" in bad.text
    unknown = login(client, "nadie", "nope")
    assert "Usuario o contraseña incorrectos" in unknown.text
    for _ in range(3):
        login(client, "ana", "nope")
    blocked = login(client, "ana", "clave-ana")          # even the right password waits
    assert blocked.status_code == 429


def test_login_rotates_session_and_logout_revokes_it(app):
    client = client_for(app)
    client.cookies.set(SESSION_COOKIE, "attacker-chosen-id", domain="testserver.local")
    login(client, "ana", "clave-ana")
    sid = client.cookies.get(SESSION_COOKIE)
    assert sid and sid != "attacker-chosen-id"
    token = csrf_of(client)
    client.post("/logout", data={"csrf_token": token}, headers=ORIGIN, follow_redirects=False)
    assert app.state.sessions.get(sid) is None


def test_session_idle_timeout():
    now = [1000.0]
    store = SessionStore(idle=60, absolute=3600, clock=lambda: now[0])
    session = store.create(1, "ana", "192.168.1.2", "iPhone")
    now[0] += 59
    assert store.get(session.sid) is not None
    now[0] += 61
    assert store.get(session.sid) is None


def test_limiter_backoff_and_lockout():
    now = [0.0]
    limiter = LoginLimiter(clock=lambda: now[0])
    for _ in range(3):
        limiter.failure("1.1.1.1", "ana")
    assert limiter.retry_after("1.1.1.1", "ana") == 0
    limiter.failure("1.1.1.1", "ana")
    assert limiter.retry_after("1.1.1.1", "ana") == 1
    for _ in range(6):
        limiter.failure("1.1.1.1", "ana")
    assert limiter.retry_after("9.9.9.9", "ANA") >= LoginLimiter.LOCK_SECONDS - 1   # per username too
    limiter.success("1.1.1.1", "ana")
    assert limiter.retry_after("1.1.1.1", "ana") == 0


# ---------------------------------------------------------------------------
# Request hardening
# ---------------------------------------------------------------------------

def test_security_headers_on_every_response(app):
    client = client_for(app)
    for path in ("/login", "/static/css/app.css"):
        headers = client.get(path).headers
        assert "default-src 'self'" in headers["content-security-policy"]
        assert "frame-ancestors 'none'" in headers["content-security-policy"]
        assert "unsafe-inline" not in headers["content-security-policy"]
        assert headers["strict-transport-security"].startswith("max-age=")
        assert headers["x-content-type-options"] == "nosniff"
        assert headers["referrer-policy"] == "same-origin"
        assert "server" not in headers
    assert logged_in(app).get("/").headers["cache-control"] == "no-store"


def test_public_addresses_are_rejected(app):
    client = TestClient(app, base_url=BASE, client=("8.8.8.8", 1234))
    assert client.get("/login").status_code == 403


def test_cross_origin_posts_are_rejected(app):
    client = client_for(app)
    no_origin = client.post("/login", data={"username": "ana", "password": "clave-ana"})
    assert no_origin.status_code == 403
    evil = client.post("/login", data={"username": "ana", "password": "clave-ana"},
                       headers={"Origin": "https://evil.example"})
    assert evil.status_code == 403
    http_origin = client.post("/login", data={"username": "ana", "password": "clave-ana"},
                              headers={"Origin": "http://testserver"})
    assert http_origin.status_code == 403


def test_null_origin_falls_back_to_referer(app):
    client = client_for(app)
    data = {"username": "ana", "password": "clave-ana"}
    assert client.post("/login", data=data, headers={"Origin": "null"}).status_code == 403
    ok = client.post("/login", data=data, follow_redirects=False,
                     headers={"Origin": "null", "Referer": f"{BASE}/login"})
    assert ok.status_code == 303


def test_writes_require_csrf_token(app, family):
    client = logged_in(app)
    url = f"/movimientos/{family['ana_txn']}/categoria"
    missing = client.post(url, data={"category_id": family["ana_cat"]}, headers=ORIGIN)
    assert missing.status_code == 403
    wrong = client.post(url, data={"category_id": family["ana_cat"]},
                        headers={**ORIGIN, "X-CSRF-Token": "x" * 40})
    assert wrong.status_code == 403
    ok = client.post(url, data={"category_id": family["ana_cat"]},
                     headers={**ORIGIN, "X-CSRF-Token": csrf_of(client)})
    assert ok.status_code == 200


def test_no_docs_or_openapi(app):
    client = logged_in(app)
    for path in ("/docs", "/redoc", "/openapi.json"):
        assert client.get(path).status_code == 404


# ---------------------------------------------------------------------------
# Per-user isolation
# ---------------------------------------------------------------------------

def test_user_never_sees_or_edits_another_users_data(app, family, factory):
    client = logged_in(app)                     # ana
    token = csrf_of(client)
    headers = {**ORIGIN, "X-CSRF-Token": token}

    assert "MERCADO BETO" not in client.get("/movimientos").text
    assert client.get(f"/movimientos/{family['beto_txn']}").status_code == 404
    assert client.post(f"/movimientos/{family['beto_txn']}/categoria",
                       data={"category_id": family["beto_cat"]}, headers=headers).status_code == 404
    # Own transaction, but another user's category id.
    assert client.post(f"/movimientos/{family['ana_txn']}/categoria",
                       data={"category_id": family["beto_cat"]}, headers=headers).status_code == 404
    # Investment operation of beto.
    assert client.get(f"/inversiones/operaciones/{family['beto_op']}").status_code == 404
    assert client.post(f"/inversiones/operaciones/{family['beto_op']}/borrar", headers=headers).status_code == 404
    assert "Acme" not in client.get("/inversiones").text

    db = factory()
    try:
        assert db.get(Transaction, family["beto_txn"]).category_id is None
        assert db.get(Transaction, family["ana_txn"]).category_id is None
        assert db.get(InvestmentTransaction, family["beto_op"]) is not None
    finally:
        db.close()


# ---------------------------------------------------------------------------
# Pages render
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("path", [
    "/", "/?rango=year", "/?rango=all", "/gastos", "/ingresos", "/gastos?periodo=anio",
    "/gastos?periodo=custom&desde=2020-01-01&hasta=2030-01-01", "/movimientos",
    "/movimientos?q=mercado&tipo=expense&categoria=sin", "/movimientos/lista?offset=0",
    "/inversiones", "/inversiones?cerradas=1", "/inversiones/evolucion", "/inversiones/distribucion",
    "/inversiones/distribucion?por=broker", "/inversiones/operaciones", "/inversiones/operaciones/nueva",
    "/inversiones/fiscal", "/mas", "/mas/categorias", "/mas/reglas", "/mas/cuentas", "/mas/importaciones",
    "/mas/categorias/nueva", "/mas/reglas/nueva", "/mas/reglas/probar?pattern=mercado&match_type=contains",
    "/importar",
    "/api/resumen", "/api/flujos/gastos", "/api/movimientos", "/api/inversiones/posiciones",
    "/api/inversiones/evolucion", "/manifest.webmanifest", "/sw.js",
])
@pytest.mark.parametrize("user", [("ana", "clave-ana"), ("beto", "clave-beto")])
def test_pages_render(app, path, user):
    client = logged_in(app, *user)
    response = client.get(path)
    assert response.status_code == 200, (path, response.text[:500])


def test_templates_have_no_inline_script_or_style(app):
    client = logged_in(app, "beto", "clave-beto")
    for path in ("/", "/gastos", "/movimientos", "/inversiones", "/inversiones/operaciones/nueva", "/mas",
                 "/mas/categorias", "/mas/reglas", "/mas/reglas/nueva", "/mas/categorias/nueva", "/importar",
                 "/mas/cuenta/contrasena", "/mas/cuenta/nombre", "/mas/cuenta/borrar", "/mas/cuenta/datos"):
        html = client.get(path).text
        assert not re.search(r"<script(?![^>]*\bsrc=)[^>]*>", html), path
        assert " style=" not in html and "<style" not in html, path
        assert not re.search(r"\son[a-z]+=", html), path
    assert "<script" not in client.get("/login").text.split("</head>")[1]
    for html in (client.post("/mas/invitaciones", headers=_hx(client)).text,
                 client_for(app).get("/registro").text.split("</head>")[1]):
        assert not re.search(r"<script", html)
        assert " style=" not in html and "<style" not in html
        assert not re.search(r"\son[a-z]+=", html)


def test_tax_xlsx_download(app):
    client = logged_in(app, "beto", "clave-beto")
    year = date.today().year
    response = client.get(f"/inversiones/fiscal/{year}.xlsx")
    assert response.status_code == 200
    assert response.content[:2] == b"PK"
    assert client.get("/inversiones/fiscal/1999.xlsx").status_code == 404


# ---------------------------------------------------------------------------
# Light edits
# ---------------------------------------------------------------------------

def test_recategorize_marks_manual_audits_and_notifies(app, family, factory):
    client = logged_in(app)
    response = client.post(f"/movimientos/{family['ana_txn']}/categoria",
                           data={"category_id": family["ana_cat"]},
                           headers={**ORIGIN, "X-CSRF-Token": csrf_of(client), "HX-Request": "true"})
    assert response.status_code == 200
    assert 'hx-swap-oob="outerHTML"' in response.text and "Supermercados" in response.text
    assert "toast" in response.headers["hx-trigger"]
    db = factory()
    try:
        txn = db.get(Transaction, family["ana_txn"])
        assert txn.category_id == family["ana_cat"]
        assert txn.category_source == CATEGORY_SOURCE_MANUAL
        log = db.query(AuditLog).filter(AuditLog.record_id == txn.id).one()
        assert "web 192.168.1.50" in log.new_value
    finally:
        db.close()
    assert app.state.changes == [(family["ana"], "transactions")]


def test_category_sheet_saves_on_tap_and_lists_every_category(app, family, factory):
    db = factory()
    try:
        tipo = Category(user_id=family["ana"], name="INGRESO", level=CategoryLevel.type,
                        transaction_type=TransactionType.income)
        db.add(tipo)
        db.flush()
        db.add(Category(user_id=family["ana"], parent_id=tipo.id, name="Nómina",
                        level=CategoryLevel.category))
        db.commit()
        expected = db.query(Category).filter(Category.user_id == family["ana"],
                                             Category.level == CategoryLevel.category).count()
    finally:
        db.close()

    client = logged_in(app)
    html = client.get(f"/movimientos/{family['ana_txn']}").text
    form = re.search(r"<form[^>]*data-cat-form[^>]*>", html).group(0)
    assert "hx-trigger=\"change[target.name=='category_id']\"" in form   # saves on tap
    assert 'type="submit"' not in html                         # no button lost below the list
    assert html.count('name="category_id"') == expected + 1    # every category + "Sin categoría"
    assert "Nómina" in html and "Supermercados" in html
    assert html.index(">GASTO <") < html.index(">INGRESO <")   # expense → GASTO first
    assert f"{expected} disponibles" in html
    assert "data-cat-search" in html


def test_investment_operation_create_validate_and_delete(app, family, factory):
    client = logged_in(app, "beto", "clave-beto")
    headers = {**ORIGIN, "X-CSRF-Token": csrf_of(client), "HX-Request": "true"}
    base = {"kind": "sell", "security_id": str(family["security"]), "trade_date": "2025-03-01",
            "broker": "DEGIRO", "price": "12", "amount_eur": "120"}

    too_many = client.post("/inversiones/operaciones", data={**base, "quantity": "50"}, headers=headers)
    assert too_many.status_code == 200 and "alert err" in too_many.text   # ledger refuses it

    ok = client.post("/inversiones/operaciones", data={**base, "quantity": "5", "amount_eur": "60"},
                     headers=headers)
    assert ok.status_code == 204
    assert "refreshPage" in ok.headers["hx-trigger"]

    db = factory()
    try:
        sale = db.query(InvestmentTransaction).filter_by(user_id=family["beto"], kind=InvestmentKind.sell).one()
        sale_id = sale.id
    finally:
        db.close()

    # Deleting the buy would leave the sale without shares → refused.
    refused = client.post(f"/inversiones/operaciones/{family['beto_op']}/borrar", headers=headers)
    assert refused.status_code == 200 and "alert err" in refused.text

    assert client.post(f"/inversiones/operaciones/{sale_id}/borrar", headers=headers).status_code == 204
    db = factory()
    try:
        assert db.get(InvestmentTransaction, sale_id) is None
        assert db.query(AuditLog).filter(AuditLog.table_affected == "investment_transactions").count() == 2
    finally:
        db.close()


def test_future_date_is_refused(app, family):
    client = logged_in(app, "beto", "clave-beto")
    response = client.post("/inversiones/operaciones", data={
        "kind": "buy", "security_id": str(family["security"]), "trade_date": "2999-01-01",
        "broker": "DEGIRO", "quantity": "1", "price": "1", "amount_eur": "1",
    }, headers={**ORIGIN, "X-CSRF-Token": csrf_of(client)})
    assert "La fecha no puede ser futura" in response.text


# ---------------------------------------------------------------------------
# Shared services
# ---------------------------------------------------------------------------

def test_assign_user_transaction_category_checks_ownership(factory, family):
    db = factory()
    try:
        with pytest.raises(LookupError):
            assign_user_transaction_category(db, family["ana"], family["beto_txn"], None)
        with pytest.raises(LookupError):
            assign_user_transaction_category(db, family["ana"], family["ana_txn"], family["beto_cat"])
        assign_user_transaction_category(db, family["ana"], family["ana_txn"], family["ana_cat"])
        assert db.get(Transaction, family["ana_txn"]).category_source == CATEGORY_SOURCE_MANUAL
    finally:
        db.close()


def test_dashboard_summary_matches_desktop_figures(factory, family):
    db = factory()
    try:
        today = date.today()
        s = dashboard_summary(db, family["ana"], date(today.year, today.month, 1), today)
        assert s.balance == 1000.0
        assert s.income == 2000.0 and s.expense == 42.5
        assert s.uncategorized == 2
        assert s.by_category == [("Sin categorizar (1)", 42.5)]
        assert round(s.savings_rate, 4) == round((2000 - 42.5) / 2000 * 100, 4)
        assert [r.issuer_name for r in s.recent][0] in ("MERCADO ANA", "NOMINA")
        assert all("BETO" not in (r.issuer_name or "") for r in s.recent)
    finally:
        db.close()


# ---------------------------------------------------------------------------
# Administración: categories and rules (W7)
# ---------------------------------------------------------------------------

def _hx(client: TestClient) -> dict:
    return {**ORIGIN, "X-CSRF-Token": csrf_of(client), "HX-Request": "true"}


def _gasto_type(factory, user_id: int) -> int:
    db = factory()
    try:
        return db.query(Category.id).filter(Category.user_id == user_id, Category.name == "GASTO").scalar()
    finally:
        db.close()


def test_category_create_rename_delete(app, family, factory):
    client = logged_in(app)
    headers = _hx(client)
    gasto = _gasto_type(factory, family["ana"])

    response = client.post("/mas/categorias", data={"parent_id": gasto, "name": "  Farmacia "}, headers=headers)
    assert response.status_code == 204 and "refreshPage" in response.headers["hx-trigger"]
    db = factory()
    try:
        farmacia = db.query(Category).filter(Category.user_id == family["ana"], Category.name == "Farmacia").one()
    finally:
        db.close()

    # Duplicates in the same group are refused with a message in the form.
    response = client.post("/mas/categorias", data={"parent_id": gasto, "name": "farmacia"}, headers=headers)
    assert response.status_code == 200 and "Ya existe" in response.text

    response = client.post(f"/mas/categorias/{farmacia.id}/renombrar", data={"name": "Salud"}, headers=headers)
    assert response.status_code == 204

    # Deleting a used category: its transactions go back to the rules
    # (uncategorized, source cleared) and the rules pointing at it go too.
    db = factory()
    try:
        txn = db.get(Transaction, family["ana_txn"])
        txn.category_id, txn.category_source = farmacia.id, CATEGORY_SOURCE_MANUAL
        db.add(AutoCategorizationRule(user_id=family["ana"], category_id=farmacia.id,
                                      match_pattern="FARMA", match_type=MatchType.contains, priority=100))
        db.commit()
    finally:
        db.close()
    sheet = client.get(f"/mas/categorias/{farmacia.id}").text
    assert "1 movimientos quedarán sin categoría" in sheet and "1 reglas" in sheet
    response = client.post(f"/mas/categorias/{farmacia.id}/borrar", headers=headers)
    assert response.status_code == 204
    db = factory()
    try:
        assert db.get(Category, farmacia.id) is None
        txn = db.get(Transaction, family["ana_txn"])
        assert txn.category_id is None and txn.category_source is None
        assert db.query(AutoCategorizationRule).filter_by(category_id=farmacia.id).count() == 0
        actions = [a.action.value for a in db.query(AuditLog).filter(AuditLog.table_affected == "categories")]
        assert actions == ["insert", "update", "delete"]
    finally:
        db.close()
    assert (family["ana"], "transactions") in app.state.changes

    # Type nodes cannot be deleted.
    response = client.post(f"/mas/categorias/{gasto}/borrar", headers=headers)
    assert response.status_code == 204 and "toastError" in response.headers["hx-trigger"]


def test_categories_and_rules_of_another_user_are_off_limits(app, family, factory):
    db = factory()
    try:
        rule = AutoCategorizationRule(user_id=family["beto"], category_id=family["beto_cat"],
                                      match_pattern="MERCADO", match_type=MatchType.contains, priority=10)
        shared = AutoCategorizationRule(user_id=None, category_id=family["beto_cat"],
                                        match_pattern="SUPER", match_type=MatchType.contains, priority=20)
        db.add_all([rule, shared])
        db.commit()
        rule_id, shared_id = rule.id, shared.id
    finally:
        db.close()

    client = logged_in(app)          # ana
    headers = _hx(client)
    beto_cat = family["beto_cat"]
    assert client.get(f"/mas/categorias/{beto_cat}").status_code == 404
    assert client.post(f"/mas/categorias/{beto_cat}/renombrar", data={"name": "X"}, headers=headers).status_code == 404
    assert client.post(f"/mas/categorias/{beto_cat}/borrar", headers=headers).status_code == 404
    assert client.get(f"/mas/reglas/{rule_id}").status_code == 404
    assert client.post(f"/mas/reglas/{rule_id}/borrar", headers=headers).status_code == 404
    assert client.post("/mas/reglas", data={"rule_id": rule_id, "category_id": family["ana_cat"],
                                            "pattern": "X"}, headers=headers).status_code == 404
    # A rule can only point at one of your own categories.
    response = client.post("/mas/reglas", data={"category_id": beto_cat, "pattern": "X"}, headers=headers)
    assert response.status_code == 200 and "categoría válida" in response.text
    # Global rules are visible but read-only on the web.
    page = client.get("/mas/reglas").text
    assert "«SUPER»" in page and "«MERCADO»" not in page
    response = client.post(f"/mas/reglas/{shared_id}/borrar", headers=headers)
    assert "toastError" in response.headers["hx-trigger"]
    db = factory()
    try:
        assert db.get(Category, beto_cat).name == "Supermercados"
        assert db.get(AutoCategorizationRule, rule_id) is not None
        assert db.get(AutoCategorizationRule, shared_id) is not None
    finally:
        db.close()


def test_rule_crud_priority_preview_and_apply(app, family, factory):
    client = logged_in(app)
    headers = _hx(client)
    response = client.post("/mas/reglas", data={"category_id": family["ana_cat"], "pattern": "mercado",
                                                "match_type": "contains", "priority": "50"}, headers=headers)
    assert response.status_code == 204
    response = client.post("/mas/reglas", data={"category_id": family["ana_cat"], "pattern": "[",
                                                "match_type": "regex"}, headers=headers)
    assert "Expresión regular no válida" in response.text
    db = factory()
    try:
        first = db.query(AutoCategorizationRule).filter_by(user_id=family["ana"]).one()
        assert first.priority == 50 and first.match_pattern == "mercado"
        second = AutoCategorizationRule(user_id=family["ana"], category_id=family["ana_cat"],
                                        match_pattern="NOMINA", match_type=MatchType.exact, priority=60)
        db.add(second)
        db.commit()
        first_id, second_id = first.id, second.id
    finally:
        db.close()

    preview = client.get("/mas/reglas/probar", params={"pattern": "MERCADO", "match_type": "contains"}).text
    assert "1 movimiento coincide" in preview and "MERCADO ANA" in preview and "BETO" not in preview

    assert client.post(f"/mas/reglas/{second_id}/prioridad", data={"direccion": "up"},
                       headers=headers).status_code == 204
    db = factory()
    try:
        assert db.get(AutoCategorizationRule, second_id).priority < db.get(AutoCategorizationRule, first_id).priority
    finally:
        db.close()

    # Apply only to uncategorized: the expense matches «mercado».
    response = client.post("/mas/reglas/aplicar", data={"modo": "sin"}, headers=headers)
    assert response.status_code == 204
    db = factory()
    try:
        txn = db.get(Transaction, family["ana_txn"])
        assert txn.category_id == family["ana_cat"] and txn.category_source == "rule"
        beto_txn = db.get(Transaction, family["beto_txn"])
        assert beto_txn.category_id is None                      # other users untouched
    finally:
        db.close()

    response = client.post(f"/mas/reglas/{first_id}/borrar", headers=headers)
    assert response.status_code == 204
    db = factory()
    try:
        assert db.get(AutoCategorizationRule, first_id) is None
    finally:
        db.close()


def test_apply_all_rules_never_overwrites_manual_categories(factory, family):
    from app.categorization.admin_service import add_category, add_rule, apply_all_rules

    db = factory()
    try:
        gasto = db.query(Category).filter_by(user_id=family["ana"], name="GASTO").one()
        other = add_category(db, family["ana"], gasto.id, "Otros")
        assign_user_transaction_category(db, family["ana"], family["ana_txn"], other.id)
        add_rule(db, family["ana"], family["ana_cat"], "MERCADO")
        apply_all_rules(db, family["ana"], overwrite=True)
        txn = db.get(Transaction, family["ana_txn"])
        assert txn.category_id == other.id and txn.category_source == CATEGORY_SOURCE_MANUAL
    finally:
        db.close()


def test_admin_writes_require_csrf(app, family):
    client = logged_in(app)
    for path, data in (("/mas/categorias", {"parent_id": "1", "name": "X"}),
                       ("/mas/reglas", {"pattern": "X"}), ("/mas/reglas/aplicar", {"modo": "sin"}),
                       ("/mas/categorias/predeterminadas", {})):
        response = client.post(path, data=data, headers=ORIGIN)
        assert response.status_code == 403, path


# ---------------------------------------------------------------------------
# Importar: CaixaBank statement upload (W8)
# ---------------------------------------------------------------------------

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def statement(tmp_path):
    from tests.test_import_pipeline import EXPORT_ROWS, _write_export

    path = tmp_path / "extracto septiembre.xlsx"
    _write_export(path, EXPORT_ROWS)
    return path.read_bytes()


def _upload(client: TestClient, content: bytes, name: str = "extracto.xlsx", **data):
    return client.post("/importar", files={"ficheros": (name, content, XLSX)},
                       data={"csrf_token": csrf_of(client), **data}, headers=ORIGIN, follow_redirects=False)


def _token(response) -> str:
    assert response.status_code == 303, response.text[:400]
    return response.headers["location"].rsplit("/", 1)[1]


def test_import_upload_preview_confirm(app, family, factory, statement):
    client = logged_in(app)
    token = _token(_upload(client, statement, "extracto septiembre.xlsx"))
    item = app.state.imports.get(token, family["ana"], next(iter(app.state.sessions.active())).sid)
    assert item is not None and item.directory.exists()

    preview = client.get(f"/importar/{token}").text
    assert "IBERDROLA CLIENTES" in preview and "extracto septiembre.xlsx" in preview
    assert "Importar <span data-import-total>5</span> movimientos" in preview
    assert 'name="importar"' not in preview                     # nothing to review on a first import

    response = client.post(f"/importar/{token}/confirmar", data={"csrf_token": csrf_of(client)},
                           headers=ORIGIN, follow_redirects=False)
    assert response.status_code == 303 and response.headers["location"].startswith("/importar/hecho?lote=")
    assert not item.directory.exists() and len(app.state.imports) == 0
    done = client.get(response.headers["location"]).text
    assert "5 movimientos importados" in done

    db = factory()
    try:
        batch = db.query(ImportBatch).filter(ImportBatch.user_id == family["ana"]).one()
        assert batch.rows_imported == 5 and batch.filename == "extracto septiembre.xlsx"
        assert db.query(Transaction).filter(Transaction.import_batch_id == batch.id,
                                            Transaction.user_id == family["ana"]).count() == 5
        log = db.query(AuditLog).filter(AuditLog.table_affected == "import_batches").one()
        assert "web 192.168.1.50" in log.new_value
    finally:
        db.close()
    assert (family["ana"], "transactions") in app.state.changes

    # The same file again: nothing new; the zero-amount row needs a decision.
    token = _token(_upload(client, statement))
    preview = client.get(f"/importar/{token}").text
    assert "¿duplicado?" in preview and 'name="importar"' in preview
    idx = re.search(r'name="importar" value="(\d+)"', preview).group(1)
    response = client.post(f"/importar/{token}/confirmar", data={"csrf_token": csrf_of(client), "importar": idx},
                           headers=ORIGIN, follow_redirects=False)
    assert response.status_code == 303
    db = factory()
    try:
        second = db.query(ImportBatch).filter(ImportBatch.user_id == family["ana"]).order_by(ImportBatch.id.desc()).first()
        assert second.rows_imported == 1 and second.rows_skipped == 4
    finally:
        db.close()


def test_import_cancel_removes_the_temp_file(app, family, statement):
    client = logged_in(app)
    token = _token(_upload(client, statement))
    item = next(iter(app.state.imports._items.values()))
    response = client.post(f"/importar/{token}/cancelar", data={"csrf_token": csrf_of(client)},
                           headers=ORIGIN, follow_redirects=False)
    assert response.status_code == 303
    assert not item.directory.exists() and len(app.state.imports) == 0
    assert client.get(f"/importar/{token}").status_code == 404


def test_import_rejects_bad_files(app, family, statement):
    client = logged_in(app)
    page = _upload(client, b"fecha;importe\n", "extracto.csv")
    assert page.status_code == 200 and "Excel (.xls o .xlsx)" in page.text
    page = _upload(client, b"not really excel", "extracto.xlsx")
    assert "no es un Excel válido" in page.text
    page = _upload(client, b"PK\x03\x04" + b"0" * (5 * 1024 * 1024), "grande.xlsx")
    assert "demasiado grande" in page.text
    page = _upload(client, b"PK\x03\x04garbage", "roto.xlsx")
    assert "No se pudo leer el extracto" in page.text
    assert len(app.state.imports) == 0
    # Bodies over the upload cap are refused before they are read: several
    # statements at once for /importar, one file for the other uploads.
    response = client.post("/importar", content=b"x" * (26 * 1024 * 1024 + 10),
                           headers={**ORIGIN, "content-type": "application/octet-stream"})
    assert response.status_code == 413
    response = client.post("/importar/inversiones", content=b"x" * (6 * 1024 * 1024 + 10),
                           headers={**ORIGIN, "content-type": "application/octet-stream"})
    assert response.status_code == 413
    # Ordinary forms have a much smaller cap.
    response = client.post("/mas/reglas", content=b"x" * (300 * 1024),
                           headers={**ORIGIN, "content-type": "application/x-www-form-urlencoded"})
    assert response.status_code == 413


def test_import_preview_is_private_to_the_uploader(app, family, factory, statement):
    ana = logged_in(app)
    token = _token(_upload(ana, statement))
    beto = logged_in(app, "beto", "clave-beto")
    assert beto.get(f"/importar/{token}").status_code == 404
    response = beto.post(f"/importar/{token}/confirmar", data={"csrf_token": csrf_of(beto)},
                         headers=ORIGIN, follow_redirects=False)
    assert response.status_code == 404
    # Another browser session of the same user cannot use it either.
    ana2 = logged_in(app)
    assert ana2.get(f"/importar/{token}").status_code == 404
    # Nor can anyone import into someone else's account.
    db = factory()
    try:
        beto_account = db.query(Account.id).filter(Account.user_id == family["beto"]).scalar()
    finally:
        db.close()
    assert _upload(ana, statement, cuenta=str(beto_account)).status_code == 404
    db = factory()
    try:
        assert db.query(ImportBatch).count() == 0
    finally:
        db.close()


def _upload_many(client: TestClient, files: list, **data):
    return client.post("/importar", files=[("ficheros", (name, content, XLSX)) for name, content in files],
                       data={"csrf_token": csrf_of(client), **data}, headers=ORIGIN, follow_redirects=False)


def test_import_several_overlapping_statements_at_once(app, family, factory, tmp_path):
    from tests.test_import_pipeline import EXPORT_ROWS, _write_export

    def export(rows, name):
        _write_export(tmp_path / name, rows)
        return (tmp_path / name).read_bytes()

    client = logged_in(app)
    september = export(EXPORT_ROWS[:4], "a.xlsx")
    overlap = export(EXPORT_ROWS[2:], "b.xlsx")
    # CaixaBank names every export alike: two uploads with the same name.
    token = _token(_upload_many(client, [("movimientos.xlsx", overlap), ("movimientos.xlsx", september),
                                         ("notas.txt", b"hola")]))
    item = app.state.imports.get(token, family["ana"], next(iter(app.state.sessions.active())).sid)
    assert item.filenames == ["movimientos (2).xlsx", "movimientos.xlsx"]       # chronological

    preview = client.get(f"/importar/{token}").text
    assert "2 extractos" in preview and "7 filas en 2 extractos" in preview
    assert "Este fichero no se importará" in preview and "notas.txt" in preview
    assert "Importar <span data-import-total>5</span> movimientos" in preview

    response = client.post(f"/importar/{token}/confirmar", data={"csrf_token": csrf_of(client)},
                           headers=ORIGIN, follow_redirects=False)
    done = client.get(response.headers["location"]).text
    assert "5 movimientos importados" in done and "2 extractos" in done
    db = factory()
    try:
        batches = db.query(ImportBatch).filter(ImportBatch.user_id == family["ana"]).order_by(ImportBatch.id).all()
        assert [(b.filename, b.rows_imported, b.rows_skipped) for b in batches] == [
            ("movimientos (2).xlsx", 4, 0), ("movimientos.xlsx", 1, 2)]
        assert db.query(Transaction).filter(Transaction.import_batch_id.in_([b.id for b in batches])).count() == 5
        assert db.query(AuditLog).filter(AuditLog.table_affected == "import_batches").count() == 2
    finally:
        db.close()


def test_import_upload_limits_and_all_files_rejected(app, family, statement, monkeypatch):
    from app.web.routers import imports

    client = logged_in(app)
    page = _upload_many(client, [(f"e{i}.xlsx", statement) for i in range(imports.MAX_FILES + 1)])
    assert page.status_code == 200 and f"como máximo {imports.MAX_FILES} extractos" in page.text
    page = _upload_many(client, [("a.csv", b"x"), ("b.xlsx", b"no excel")])
    assert "Ninguno de los extractos se puede importar" in page.text
    assert "a.csv" in page.text and "b.xlsx" in page.text
    monkeypatch.setattr(imports, "MAX_TOTAL_BYTES", len(statement) + 1)
    page = _upload_many(client, [("a.xlsx", statement), ("b.xlsx", statement)])
    assert "Súbelos en varias veces" in page.text
    assert len(app.state.imports) == 0


def test_import_done_page_only_shows_your_own_batches(app, family, factory, statement):
    ana = logged_in(app)
    token = _token(_upload(ana, statement))
    location = ana.post(f"/importar/{token}/confirmar", data={"csrf_token": csrf_of(ana)},
                        headers=ORIGIN, follow_redirects=False).headers["location"]
    beto = logged_in(app, "beto", "clave-beto")
    assert beto.get(location).status_code == 404
    assert ana.get(location + "&lote=999").status_code == 404
    assert ana.get("/importar/hecho").status_code == 404


def test_import_upload_requires_csrf(app, family, statement):
    client = logged_in(app)
    response = client.post("/importar", files={"ficheros": ("e.xlsx", statement, XLSX)},
                           headers=ORIGIN, follow_redirects=False)
    assert response.status_code == 403
    assert len(app.state.imports) == 0


def test_orphaned_upload_folders_are_swept(tmp_path, monkeypatch):
    import os
    import time

    from app.web import uploads

    monkeypatch.setattr(uploads.tempfile, "gettempdir", lambda: str(tmp_path))
    old = tmp_path / (uploads.PREFIX + "old")
    fresh = tmp_path / (uploads.PREFIX + "fresh")
    for d in (old, fresh):
        d.mkdir()
        (d / "extracto.xlsx").write_bytes(b"PK")
    past = time.time() - uploads.TTL - 60
    os.utime(old, (past, past))
    assert uploads.sweep_orphans() == 1
    assert not old.exists() and fresh.exists()


def test_apply_category_to_similar_movements(app, family, factory):
    db = factory()
    try:
        account_id = db.get(Transaction, family["ana_txn"]).account_id
        twins = [Transaction(user_id=family["ana"], account_id=account_id, transaction_date=date.today(),
                             amount=-3.0, type=TransactionType.expense, issuer_name="MERCADO ANA")
                 for _ in range(2)]
        db.add_all(twins)
        db.commit()
        twin_ids = [t.id for t in twins]
    finally:
        db.close()

    client = logged_in(app)
    sheet = client.get(f"/movimientos/{family['ana_txn']}").text
    assert 'name="similares"' in sheet and "Aplicar también a 2 movimientos iguales" in sheet
    assert "change[target.name=='category_id']" in sheet     # ticking the box alone does not save

    response = client.post(f"/movimientos/{family['ana_txn']}/categoria",
                           data={"category_id": family["ana_cat"], "similares": "1"},
                           headers={**ORIGIN, "X-CSRF-Token": csrf_of(client), "HX-Request": "true"})
    assert response.status_code == 200
    assert response.text.count('hx-swap-oob="outerHTML"') == 3
    assert "3 movimientos" in response.headers["hx-trigger"]
    db = factory()
    try:
        for txn_id in twin_ids:
            txn = db.get(Transaction, txn_id)
            assert (txn.category_id, txn.category_source) == (family["ana_cat"], CATEGORY_SOURCE_MANUAL)
        # beto's "MERCADO BETO" is a different merchant and another user: untouched.
        assert db.get(Transaction, family["beto_txn"]).category_id is None
    finally:
        db.close()
    # Nothing left to apply to: the box is no longer offered.
    assert 'name="similares"' not in client.get(f"/movimientos/{family['ana_txn']}").text


def test_category_sheet_can_create_a_rule(app, family, factory):
    db = factory()
    try:
        account_id = db.get(Transaction, family["ana_txn"]).account_id
        twin = Transaction(user_id=family["ana"], account_id=account_id, transaction_date=date.today(),
                           amount=-3.0, type=TransactionType.expense, issuer_name="MERCADO ANA")
        db.add(twin)
        db.commit()
        twin_id = twin.id
    finally:
        db.close()

    client = logged_in(app)
    assert 'name="regla"' in client.get(f"/movimientos/{family['ana_txn']}").text
    response = client.post(f"/movimientos/{family['ana_txn']}/categoria",
                           data={"category_id": family["ana_cat"], "regla": "1"},
                           headers={**ORIGIN, "X-CSRF-Token": csrf_of(client), "HX-Request": "true"})
    assert response.status_code == 200
    assert "regla creada" in response.headers["hx-trigger"]
    db = factory()
    try:
        rule = db.query(AutoCategorizationRule).one()
        assert (rule.user_id, rule.match_pattern, rule.category_id) == (family["ana"], "MERCADO ANA", family["ana_cat"])
        assert db.get(Transaction, twin_id).category_id == family["ana_cat"]
        assert db.query(AuditLog).filter(AuditLog.table_affected == "auto_categorization_rules").count() == 1
    finally:
        db.close()
    # Covered by a rule now: the option is gone.
    assert 'name="regla"' not in client.get(f"/movimientos/{twin_id}").text


# ---------------------------------------------------------------------------
# Categorization library: export / import
# ---------------------------------------------------------------------------

def _library_upload(client: TestClient, content: bytes, name: str = "cats.profin-cats.json"):
    return client.post("/importar/categorias", files={"fichero": (name, content, "application/json")},
                       data={"csrf_token": csrf_of(client)}, headers=ORIGIN, follow_redirects=False)


def _rules_of(factory, user_id):
    from app.db.models import AutoCategorizationRule
    db = factory()
    try:
        return sorted(r.match_pattern for r in db.query(AutoCategorizationRule).filter_by(user_id=user_id))
    finally:
        db.close()


def _add_web_rule(factory, user_id, category_id, pattern, global_=False):
    from app.db.models import AutoCategorizationRule
    db = factory()
    db.add(AutoCategorizationRule(user_id=None if global_ else user_id, category_id=category_id,
                                  match_pattern=pattern))
    db.commit()
    db.close()


def test_library_export_downloads_only_my_personal_rules(app, family, factory):
    import json
    _add_web_rule(factory, family["ana"], family["ana_cat"], "MERCADONA")
    _add_web_rule(factory, family["beto"], family["beto_cat"], "SOLO BETO")
    _add_web_rule(factory, family["ana"], family["ana_cat"], "GLOBAL", global_=True)
    client = logged_in(app)
    assert "Exportar / importar" in client.get("/mas/categorias").text
    # The body is hx-boost'ed: without opting out, htmx swaps the JSON into
    # the page instead of letting the browser save the file.
    page = client.get("/mas/categorias/biblioteca").text
    assert 'href="/mas/categorias/exportar" hx-boost="false" download' in page
    response = client.get("/mas/categorias/exportar")
    assert response.status_code == 200
    assert "attachment" in response.headers["content-disposition"]
    assert ".profin-cats.json" in response.headers["content-disposition"]
    payload = json.loads(response.content)
    assert [r["pattern"] for r in payload["rules"]] == ["MERCADONA"]
    assert payload["source_user"] == "ana"


def test_library_upload_preview_confirm(app, family, factory):
    _add_web_rule(factory, family["ana"], family["ana_cat"], "MERCADONA")
    ana_client = logged_in(app)
    exported = ana_client.get("/mas/categorias/exportar").content

    client = logged_in(app, "beto", "clave-beto")
    _add_web_rule(factory, family["beto"], family["beto_cat"], "ALDI")
    token = _token(_library_upload(client, exported))
    preview = client.get(f"/importar/categorias/{token}")
    assert preview.status_code == 200 and "MERCADONA" in preview.text and "Fusionar" in preview.text

    # Replace without the explicit tick goes back to the preview.
    response = client.post(f"/importar/categorias/{token}/confirmar",
                           data={"csrf_token": csrf_of(client), "modo": "replace"},
                           headers=ORIGIN, follow_redirects=False)
    assert response.status_code == 303 and "modo=replace" in response.headers["location"]
    assert _rules_of(factory, family["beto"]) == ["ALDI"]

    response = client.post(f"/importar/categorias/{token}/confirmar",
                           data={"csrf_token": csrf_of(client), "modo": "merge"}, headers=ORIGIN)
    assert response.status_code == 200 and "importadas" in response.text
    assert _rules_of(factory, family["beto"]) == ["ALDI", "MERCADONA"]
    assert (family["beto"], "transactions") in app.state.changes
    # The token is single-use.
    assert client.get(f"/importar/categorias/{token}").status_code == 404

    token = _token(_library_upload(client, exported))
    response = client.post(f"/importar/categorias/{token}/confirmar",
                           data={"csrf_token": csrf_of(client), "modo": "replace", "entiendo": "1"},
                           headers=ORIGIN)
    assert response.status_code == 200
    assert _rules_of(factory, family["beto"]) == ["MERCADONA"]


def test_library_web_never_imports_global_rules(app, family, factory):
    from app.categorization import library_io
    from app.db.models import AutoCategorizationRule
    payload = {
        "format": library_io.FORMAT_NAME, "format_version": 1,
        "types": [{"name": "GASTO", "transaction_type": "expense", "categories": ["Supermercados"]}],
        "rules": [{"type": "GASTO", "category": "Supermercados", "pattern": "CARREFOUR",
                   "match_type": "contains", "priority": 100, "scope": "global"}],
    }
    client = logged_in(app)
    token = _token(_library_upload(client, library_io.dumps(payload)))
    preview = client.get(f"/importar/categorias/{token}").text
    assert "reglas globales" in preview and "No hay nada que importar" in preview
    client.post(f"/importar/categorias/{token}/confirmar",
                data={"csrf_token": csrf_of(client), "modo": "merge"}, headers=ORIGIN)
    db = factory()
    try:
        assert db.query(AutoCategorizationRule).filter_by(user_id=None).count() == 0
    finally:
        db.close()


def test_library_upload_rejects_bad_files_and_needs_csrf(app, family):
    client = logged_in(app)
    response = _library_upload(client, b"esto no es json")
    assert response.status_code == 400 and "JSON" in response.text
    response = client.post("/importar/categorias", files={"fichero": ("c.json", b"{}", "application/json")},
                           headers=ORIGIN, follow_redirects=False)
    assert response.status_code == 403


def test_library_preview_is_private_to_the_uploader(app, family, factory):
    from app.categorization import library_io
    payload = {"format": library_io.FORMAT_NAME, "format_version": 1,
               "types": [{"name": "GASTO", "transaction_type": "expense", "categories": ["Nueva"]}],
               "rules": []}
    token = _token(_library_upload(logged_in(app), library_io.dumps(payload)))
    other = logged_in(app, "beto", "clave-beto")
    assert other.get(f"/importar/categorias/{token}").status_code == 404
    response = other.post(f"/importar/categorias/{token}/confirmar",
                          data={"csrf_token": csrf_of(other)}, headers=ORIGIN)
    assert response.status_code == 404


# ---------------------------------------------------------------------------
# User manual (/ayuda) and the DEMO account
# ---------------------------------------------------------------------------
def test_manual_is_public_with_its_own_csp(app):
    client = client_for(app)
    assert client.get("/ayuda", follow_redirects=False).headers["location"] == "/ayuda/"
    response = client.get("/ayuda/")
    assert response.status_code == 200
    assert "Manual de FAMILYFIN" in response.text
    assert '<a class="brand" href="/"' in response.text      # back to the app
    csp = response.headers["content-security-policy"]
    assert "script-src 'sha256-" in csp and "unsafe-inline" not in csp.split("script-src")[1].split(";")[0]
    assert client.get("/ayuda/img/w_resumen.png").headers["content-type"] == "image/png"
    assert client.get("/ayuda/manual.pdf").headers["content-type"] == "application/pdf"


def test_manual_images_reject_other_files(app):
    client = client_for(app)
    for name in ("..%2Findex.html", "..%2F..%2F..%2Fapp%2Fconfig.py", "nope.png", "index.html"):
        assert client.get(f"/ayuda/img/{name}").status_code == 404


def test_manual_is_linked_without_htmx_boost(app):
    link = 'href="/ayuda/" hx-boost="false"'
    assert "/ayuda/" in client_for(app).get("/login").text     # the login page is not boosted
    client = logged_in(app)
    assert link in client.get("/").text                        # sidebar
    assert link in client.get("/mas").text


def test_demo_account_does_not_see_family_global_rules(app, family, factory):
    from app.config import DEMO_USERNAME

    db = factory()
    demo = create_user(db, DEMO_USERNAME, "123456", enforce_policy=False)
    db.close()
    _add_web_rule(factory, demo.id, family["ana_cat"], "PATRON DEMO")
    _add_web_rule(factory, family["ana"], family["ana_cat"], "TIENDA DE LA FAMILIA", global_=True)

    ana = logged_in(app)
    assert "TIENDA DE LA FAMILIA" in ana.get("/mas/reglas").text
    visitor = logged_in(app, DEMO_USERNAME, "123456")
    page = visitor.get("/mas/reglas").text
    assert "TIENDA DE LA FAMILIA" not in page and "PATRON DEMO" in page
    reglas = re.search(r'href="/mas/reglas">.*?<div class="v num">(\d+)</div>', visitor.get("/mas").text, re.S)
    assert reglas.group(1) == "1"


# ---------------------------------------------------------------------------
# Inversiones: DEGIRO statement upload (W9d)
# ---------------------------------------------------------------------------

def _degiro_xlsx(tmp_path, rows=None, name="Account.xlsx") -> bytes:
    from tests.test_investments_import import ROWS, write_xlsx

    return write_xlsx(tmp_path / name, ROWS if rows is None else rows).read_bytes()


def _broker_upload(client: TestClient, content: bytes, name: str = "Account.xlsx", csrf: bool = True):
    data = {"csrf_token": csrf_of(client)} if csrf else {}
    return client.post("/importar/inversiones", files={"fichero": (name, content, XLSX)},
                       data=data, headers=ORIGIN, follow_redirects=False)


def _wait_job(item):
    if item.job is not None:
        item.job.exception(timeout=10)              # background step finished (or failed)


def _broker_item(app, response):
    assert response.status_code == 303, response.text[:400]
    token = response.headers["location"].rsplit("/", 1)[1]
    item = app.state.investment_imports._items[token]
    _wait_job(item)
    return token, item


# Preview order is by product name: Acme, Siemens, US Co, World ETF.
BROKER_SYMBOLS = ["ACM.DE", "SIX.DE", "USC", "WLD.AS"]


def _broker_confirm(app, client, token, symbols):
    """Confirm; when new securities need data, wait for it and poll like the page does.
    Returns the poll's HX-Redirect target, or the confirm response."""
    response = client.post(f"/importar/inversiones/{token}/confirmar",
                           data={"csrf_token": csrf_of(client), "simbolo": symbols},
                           headers=ORIGIN, follow_redirects=False)
    item = app.state.investment_imports._items.get(token)
    if response.status_code == 303 and item is not None and item.job is not None:
        _wait_job(item)
        poll = client.post(f"/importar/inversiones/{token}/estado", headers=_hx(client))
        assert poll.status_code == 204, poll.text[:300]
        return poll.headers["HX-Redirect"]
    return response


def _degiro_ops(factory, user_id):
    db = factory()
    try:
        return db.query(InvestmentTransaction).filter(InvestmentTransaction.user_id == user_id,
                                                      InvestmentTransaction.broker == "DEGIRO").all()
    finally:
        db.close()


@pytest.fixture
def market(app):
    from tests.test_investments_import import FakeProvider

    app.state.market_provider = FakeProvider()
    return app.state.market_provider


def test_broker_import_preview_confirm_and_overlap_without_duplicates(app, family, factory, market, tmp_path):
    from tests.test_investments_import import ROWS

    client = logged_in(app)
    assert "/inversiones/importar" in client.get("/inversiones").text        # entry point

    # 1. An older statement (everything but the last sale).
    token, item = _broker_item(app, _broker_upload(client, _degiro_xlsx(tmp_path, ROWS[:-2], "old.xlsx")))
    page = client.get(f"/importar/inversiones/{token}").text
    assert "Acme Corp." in page and 'value="ACM.DE"' in page                 # proposed symbol
    assert "Sin coincidencia" in page                                         # Siemens: type it
    assert page.count('name="simbolo"') == 4
    target = _broker_confirm(app, client, token, BROKER_SYMBOLS)
    assert target == f"/importar/inversiones/{token}/hecho"
    assert "7 operaciones importadas" in client.get(target).text
    assert len(_degiro_ops(factory, family["ana"])) == 7
    assert (family["ana"], "investments") in app.state.changes

    # 2. The full statement overlaps: only the sale is new.
    token, item = _broker_item(app, _broker_upload(client, _degiro_xlsx(tmp_path)))
    page = client.get(f"/importar/inversiones/{token}").text
    assert page.count('name="simbolo"') == 1 and "Ya está en el catálogo" in page
    assert "Importar 1 operaciones" in page
    response = _broker_confirm(app, client, token, ["WLD.AS"])                 # known symbol: no fetch
    assert response.status_code == 303 and response.headers["location"].endswith("/hecho")
    assert "1 operaciones importadas" in client.get(response.headers["location"]).text
    refs = [t.broker_ref for t in _degiro_ops(factory, family["ana"])]
    assert len(refs) == 8 and len(set(refs)) == 8

    # 3. The same file again: nothing to do, and confirming anyway changes nothing.
    token, item = _broker_item(app, _broker_upload(client, _degiro_xlsx(tmp_path)))
    page = client.get(f"/importar/inversiones/{token}").text
    assert "Nada nuevo que importar" in page and 'name="simbolo"' not in page
    response = client.post(f"/importar/inversiones/{token}/confirmar", data={"csrf_token": csrf_of(client)},
                           headers=ORIGIN, follow_redirects=False)
    assert response.status_code == 303
    assert len(_degiro_ops(factory, family["ana"])) == 8
    db = factory()
    try:
        logs = db.query(AuditLog).filter(AuditLog.table_affected == "investment_transactions").all()
        assert len(logs) == 2 and all("web 192.168.1.50" in log.new_value for log in logs)
    finally:
        db.close()


def test_broker_import_cannot_be_confirmed_twice(app, family, factory, market, tmp_path):
    client = logged_in(app)
    token, item = _broker_item(app, _broker_upload(client, _degiro_xlsx(tmp_path)))
    for _ in range(2):                                    # double tap
        response = client.post(f"/importar/inversiones/{token}/confirmar",
                               data={"csrf_token": csrf_of(client), "simbolo": BROKER_SYMBOLS},
                               headers=ORIGIN, follow_redirects=False)
        assert response.status_code == 303
    _wait_job(item)
    for _ in range(2):
        client.post(f"/importar/inversiones/{token}/estado", headers=_hx(client))
    assert len(_degiro_ops(factory, family["ana"])) == 8


def test_broker_import_skips_operations_imported_meanwhile(app, family, factory, market, tmp_path):
    from app.investments.importers.degiro import parse_statement
    from app.investments.importers.service import apply_import, fetch_security_infos, symbols_to_fetch
    from tests.test_investments_import import ACME, ETF, SIE, USCO, write_xlsx

    client = logged_in(app)
    token, item = _broker_item(app, _broker_upload(client, _degiro_xlsx(tmp_path)))
    # The PC app imports the same statement while the preview is open.
    symbols = dict(zip([ACME, SIE, USCO, ETF], BROKER_SYMBOLS))
    db = factory()
    try:
        statement = parse_statement(write_xlsx(tmp_path / "pc.xlsx"))
        infos = fetch_security_infos(market, symbols_to_fetch(db, symbols.values()))
        apply_import(db, family["ana"], statement, symbols, infos)
    finally:
        db.close()
    response = _broker_confirm(app, client, token, BROKER_SYMBOLS)
    assert response.status_code == 303
    assert "0 operaciones importadas" in client.get(response.headers["location"]).text
    assert len(_degiro_ops(factory, family["ana"])) == 8


def test_broker_import_is_all_or_nothing_and_keeps_the_preview(app, family, factory, market, tmp_path):
    from tests.test_investments_import import ROWS

    client = logged_in(app)
    rows = [r for r in ROWS if r[-1] != "o2"]                 # the ETF sale has no purchase
    token, item = _broker_item(app, _broker_upload(client, _degiro_xlsx(tmp_path, rows)))
    target = _broker_confirm(app, client, token, BROKER_SYMBOLS)
    assert target == f"/importar/inversiones/{token}"
    page = client.get(target).text
    assert "No se ha importado nada" in page and 'name="simbolo"' in page
    assert _degiro_ops(factory, family["ana"]) == []
    db = factory()
    try:
        assert db.query(Security).count() == 1               # new securities rolled back too
    finally:
        db.close()


def test_broker_import_validates_symbols(app, family, factory, market, tmp_path):
    client = logged_in(app)
    token, item = _broker_item(app, _broker_upload(client, _degiro_xlsx(tmp_path)))
    response = _broker_confirm(app, client, token, ["ACM.DE", "", "<script>", "WLD.AS"])
    assert response.status_code == 400
    assert "Falta el símbolo de «Siemens X AG»" in response.text and "no es válido" in response.text
    response = _broker_confirm(app, client, token, ["ACM.DE"])            # wrong field count
    assert response.status_code == 400
    assert _degiro_ops(factory, family["ana"]) == []


def test_broker_import_reports_market_data_errors(app, family, factory, market, tmp_path):
    from app.investments.market_data import MarketDataError

    def broken(symbol):
        raise MarketDataError("Yahoo no responde")

    market.security_info = broken
    client = logged_in(app)
    token, item = _broker_item(app, _broker_upload(client, _degiro_xlsx(tmp_path)))
    target = _broker_confirm(app, client, token, BROKER_SYMBOLS)
    page = client.get(target).text
    assert "Yahoo no responde" in page and 'name="simbolo"' in page
    assert _degiro_ops(factory, family["ana"]) == []


def test_broker_import_is_private_to_the_uploader(app, family, factory, market, tmp_path):
    ana = logged_in(app)
    token, item = _broker_item(app, _broker_upload(ana, _degiro_xlsx(tmp_path)))
    beto = logged_in(app, "beto", "clave-beto")
    assert beto.get(f"/importar/inversiones/{token}").status_code == 404
    response = beto.post(f"/importar/inversiones/{token}/confirmar",
                         data={"csrf_token": csrf_of(beto), "simbolo": BROKER_SYMBOLS},
                         headers=ORIGIN, follow_redirects=False)
    assert response.status_code == 404
    assert logged_in(app).get(f"/importar/inversiones/{token}").status_code == 404   # other session
    assert _degiro_ops(factory, family["ana"]) == [] and len(_degiro_ops(factory, family["beto"])) == 1


def test_broker_import_rejects_bad_files(app, family, market, tmp_path):
    from tests.test_investments_import_myinvestor import write_csv as myinvestor_csv

    client = logged_in(app)
    page = _broker_upload(client, b"%PDF-1.4", "extracto.pdf")
    assert page.status_code == 400 and "Excel (.xlsx) o CSV" in page.text
    page = _broker_upload(client, b"not really excel", "Account.xlsx")
    assert "no es un Excel válido" in page.text
    page = _broker_upload(client, "Fecha;Importe\n\xe9".encode("latin-1"), "Account.csv")
    assert "UTF-8" in page.text
    page = _broker_upload(client, b"Concepto,Importe\nx,1\n", "Account.csv")
    assert "No se reconoce el archivo" in page.text
    page = _broker_upload(client, myinvestor_csv(tmp_path / "mi.csv").read_bytes(), "mi.csv")
    assert "desde la aplicación del PC" in page.text
    assert len(app.state.investment_imports) == 0
    response = _broker_upload(client, _degiro_xlsx(tmp_path), csrf=False)
    assert response.status_code == 403 and len(app.state.investment_imports) == 0


def test_broker_import_cancel(app, family, market, tmp_path):
    client = logged_in(app)
    token, item = _broker_item(app, _broker_upload(client, _degiro_xlsx(tmp_path)))
    response = client.post(f"/importar/inversiones/{token}/cancelar", data={"csrf_token": csrf_of(client)},
                           headers=ORIGIN, follow_redirects=False)
    assert response.status_code == 303 and len(app.state.investment_imports) == 0
    assert client.get(f"/importar/inversiones/{token}").status_code == 404


def test_broker_import_starts_a_price_refresh(app, family, factory, market, tmp_path):
    client = logged_in(app)
    token, item = _broker_item(app, _broker_upload(client, _degiro_xlsx(tmp_path)))
    _broker_confirm(app, client, token, BROKER_SYMBOLS)
    result = app.state.price_refresh.wait(timeout=10)
    assert result is not None and result.ok and "ACM.DE" in result.updated


# ---------------------------------------------------------------------------
# Inversiones: «Actualizar cotizaciones» from the web (W9b)
# ---------------------------------------------------------------------------

class PricesProvider(OfflineProvider):
    """Closes for ACME (held by beto in the fixture); optionally blocks or fails."""

    def __init__(self, gate=None, error=None):
        self.calls, self.gate, self.error = 0, gate, error

    def daily_closes(self, symbols, start, end=None):
        from datetime import timedelta

        self.calls += 1
        if self.gate is not None:
            assert self.gate.wait(timeout=10)
        if self.error:
            from app.investments.market_data import MarketDataError
            raise MarketDataError(self.error)
        end = end or date.today()
        days = [start + timedelta(days=i) for i in range((end - start).days + 1)]
        return {"ACME": [(d, 12.5) for d in days if d.weekday() < 5]} if "ACME" in symbols else {}


def _refresh(client, htmx=True):
    headers = _hx(client) if htmx else {**ORIGIN}
    data = {} if htmx else {"csrf_token": csrf_of(client)}
    return client.post("/inversiones/cotizaciones/actualizar", data=data, headers=headers,
                       follow_redirects=False)


def test_price_refresh_runs_on_the_server_and_reloads_the_page(app, family, factory):
    from app.db.models import PriceHistory

    provider = app.state.market_provider = PricesProvider()
    client = logged_in(app, "beto", "clave-beto")
    page = client.get("/inversiones").text
    assert 'action="/inversiones/cotizaciones/actualizar"' in page and 'id="price-status"' in page

    response = _refresh(client)
    assert response.status_code == 200 and 'id="price-status"' in response.text
    result = app.state.price_refresh.wait(timeout=10)
    assert result.ok and result.updated == ["ACME"] and provider.calls == 1

    status = client.get("/inversiones/cotizaciones/estado", headers=_hx(client))
    assert "Cotizaciones actualizadas a las" in status.text and "hx-get" not in status.text
    assert "refreshPage" in status.headers["HX-Trigger"]
    assert "Cotizaciones actualizadas a las" in client.get("/inversiones").text
    db = factory()
    try:
        assert db.query(PriceHistory).filter(PriceHistory.security_id == family["security"]).count() > 0
        log = db.query(AuditLog).filter(AuditLog.table_affected == "price_history").one()
        assert log.user_id == family["beto"] and "web 192.168.1.50" in log.new_value
    finally:
        db.close()
    assert (family["beto"], "investments") in app.state.changes

    # A second tap right after: no new download.
    again = _refresh(client)
    assert "se acaban de actualizar" in again.text and provider.calls == 1


def test_price_refresh_is_one_at_a_time(app, family):
    import threading

    gate = threading.Event()
    provider = app.state.market_provider = PricesProvider(gate=gate)
    ana, beto = logged_in(app), logged_in(app, "beto", "clave-beto")
    try:
        first = _refresh(beto)
        assert "Actualizando las cotizaciones" in first.text and 'hx-get="/inversiones/cotizaciones/estado"' in first.text
        second = _refresh(ana)                                   # other user, same job
        assert "Actualizando las cotizaciones" in second.text
        poll = ana.get("/inversiones/cotizaciones/estado", headers=_hx(ana))
        assert "hx-get" in poll.text and "HX-Trigger" not in poll.headers
        assert "disabled" in ana.get("/inversiones").text.split('aria-label="Actualizar cotizaciones"')[1][:80]
    finally:
        gate.set()
    app.state.price_refresh.wait(timeout=10)
    assert provider.calls == 1


def test_price_refresh_waits_for_a_desktop_refresh(app, family):
    from app.investments.prices import REFRESH_LOCK

    provider = app.state.market_provider = PricesProvider()
    client = logged_in(app, "beto", "clave-beto")
    with REFRESH_LOCK:                                   # the PC app is downloading
        _refresh(client)
        assert app.state.price_refresh.running and provider.calls == 0
    assert app.state.price_refresh.wait(timeout=10).ok and provider.calls == 1


def test_price_refresh_reports_errors_and_can_be_retried(app, family):
    provider = app.state.market_provider = PricesProvider(error="Yahoo no responde")
    client = logged_in(app, "beto", "clave-beto")
    _refresh(client)
    assert app.state.price_refresh.wait(timeout=10).error == "Yahoo no responde"
    status = client.get("/inversiones/cotizaciones/estado", headers=_hx(client))
    assert "Yahoo no responde" in status.text and "toastError" in status.headers["HX-Trigger"]
    provider.error = None
    _refresh(client)                                     # a failure does not block a retry
    assert app.state.price_refresh.wait(timeout=10).ok and provider.calls == 2


def test_price_refresh_without_htmx_and_without_csrf(app, family):
    app.state.market_provider = PricesProvider()
    client = logged_in(app, "beto", "clave-beto")
    response = client.post("/inversiones/cotizaciones/actualizar", headers=ORIGIN, follow_redirects=False)
    assert response.status_code == 403 and not app.state.price_refresh.running
    response = _refresh(client, htmx=False)
    assert response.status_code == 303 and response.headers["location"] == "/inversiones"
    assert app.state.price_refresh.wait(timeout=10).ok


# ---------------------------------------------------------------------------
# Inversiones: broker cash in the portfolio value
# ---------------------------------------------------------------------------

def test_broker_import_stores_the_cash_and_the_portfolio_shows_the_total(app, family, factory, market, tmp_path):
    from tests.test_investments_cash import write_with_balances
    from tests.test_investments_import import ROWS

    def statement(final_balance, name):
        balances = [("EUR", 10.0)] * (len(ROWS) - 1) + [("EUR", final_balance)]
        return write_with_balances(tmp_path / name, ROWS, balances).read_bytes()

    client = logged_in(app)
    token, item = _broker_item(app, _broker_upload(client, statement(1002.42, "oct.xlsx")))
    page = client.get(f"/importar/inversiones/{token}").text
    assert "Efectivo en DEGIRO según el extracto" in page and "1.002,42 EUR" in page
    target = _broker_confirm(app, client, token, BROKER_SYMBOLS)
    done = client.get(target).text
    assert "8 operaciones importadas" in done and "Ya cuenta en el valor total" in done

    page = client.get("/inversiones").text
    assert "Valor total · valores + efectivo" in page and "Efectivo" in page
    assert "según el extracto del 03/03/2026" in page
    assert "Efectivo" in client.get("/inversiones/distribucion").text
    api = client.get("/api/inversiones/posiciones").json()["totals"]
    assert api["cash_eur"] == pytest.approx(1002.42)
    assert api["total_value_eur"] == pytest.approx(api["market_value_eur"] + 1002.42)

    # Same operations, new balance: only the cash changes, and it can be saved.
    token, item = _broker_item(app, _broker_upload(client, statement(50.0, "oct2.xlsx")))
    page = client.get(f"/importar/inversiones/{token}").text
    assert "Actualizar el efectivo" in page and "antes 1.002,42" in page and 'name="simbolo"' not in page
    response = client.post(f"/importar/inversiones/{token}/confirmar", data={"csrf_token": csrf_of(client)},
                           headers=ORIGIN, follow_redirects=False)
    assert response.status_code == 303 and response.headers["location"].endswith("/hecho")
    assert "Efectivo actualizado" in client.get(response.headers["location"]).text
    assert client.get("/api/inversiones/posiciones").json()["totals"]["cash_eur"] == pytest.approx(50.0)
    assert len(_degiro_ops(factory, family["ana"])) == 8
    # Nobody else sees it.
    beto = logged_in(app, "beto", "clave-beto")
    assert beto.get("/api/inversiones/posiciones").json()["totals"]["cash_eur"] == 0


def test_app_version_is_shown_only_after_login(app):
    from app.config import APP_VERSION

    assert APP_VERSION not in client_for(app).get("/login").text      # not disclosed before sign-in
    client = logged_in(app)
    assert f'<div class="ver num" title="Versión de FAMILYFIN">v{APP_VERSION}</div>' in client.get("/").text
    assert f"FAMILYFIN versión {APP_VERSION}" in client.get("/mas").text


# ---------------------------------------------------------------------------
# Ajustes › Familia: invitations; /registro: joining with a code (W10)
# ---------------------------------------------------------------------------

CODE_RE = re.compile(r'data-invite-text>([2-9A-Z]{4}-[2-9A-Z]{4})<')


def _invite(client: TestClient) -> str:
    response = client.post("/mas/invitaciones", headers=_hx(client))
    assert response.status_code == 200, response.text
    return CODE_RE.search(response.text).group(1)


def _register(client: TestClient, code: str, username: str = "carla", password: str = "clave de carla",
              confirm=None):
    return client.post("/registro", data={"codigo": code, "username": username, "password": password,
                                          "confirm": password if confirm is None else confirm},
                       headers=ORIGIN, follow_redirects=False)


def test_member_invites_and_the_new_person_joins_signed_in(app, family, factory, caplog):
    from app.db.models import User

    caplog.set_level("DEBUG")
    ana = logged_in(app)
    page = ana.get("/mas").text
    assert "Familia" in page and ">beto<" in page and 'hx-post="/mas/invitaciones"' in page
    assert "/mas/usuarios" not in page                       # no direct creation any more

    code = _invite(ana)
    assert "Invitación pendiente" in ana.get("/mas").text

    guest = client_for(app)
    login_page = guest.get("/login").text
    assert 'href="/registro"' in login_page and "¿Tienes un código de invitación?" in login_page
    step2 = guest.post("/registro/codigo", data={"codigo": code.lower().replace("-", " ")}, headers=ORIGIN)
    assert step2.status_code == 200 and "Te invita <b>ana</b>" in step2.text
    assert f'name="codigo" value="{code}"' in step2.text
    hint = guest.get("/registro/disponible", params={"username": "Beto", "codigo": code})
    assert "Ya existe el usuario «beto»" in hint.text

    response = _register(guest, code, username=" Carla ")
    assert response.status_code == 303 and response.headers["location"] == "/mas?bienvenida=1"
    welcome = guest.get("/mas?bienvenida=1").text
    assert "Bienvenido/a, Carla" in welcome and "Cargar categorías por defecto" in welcome
    assert any(area == "user_created" for _, area in app.state.changes)

    db = factory()
    try:
        carla = db.query(User).filter(User.username == "Carla").one()
        assert db.query(Category).filter(Category.user_id == carla.id).count() == 0      # starts empty
        entry = db.query(AuditLog).filter(AuditLog.table_affected == "users").one()
        assert entry.user_id == carla.id and "invitación de «ana»" in entry.new_value
        issued = db.query(AuditLog).filter(AuditLog.table_affected == "invitations").one()
        assert issued.user_id == family["ana"]
        for row in db.query(AuditLog):
            assert code not in f"{row.old_value} {row.new_value}"
            assert code.replace("-", "") not in f"{row.old_value} {row.new_value}"
    finally:
        db.close()
    # The app's own logs (profin.log) never carry the code. (The test client's
    # httpx logger prints request URLs; the real server runs without access log.)
    app_logs = " | ".join(r.getMessage() for r in caplog.records if r.name.startswith("app."))
    assert "invitación" in app_logs
    assert code not in app_logs and code.replace("-", "") not in app_logs
    # Single use: the same code does not work again.
    again = client_for(app).post("/registro/codigo", data={"codigo": code}, headers=ORIGIN)
    assert again.status_code == 400 and "no válido o caducado" in again.text
    assert "Invitación pendiente" not in ana.get("/mas").text


def test_registration_validates_without_spending_the_invitation(app, family, factory):
    from app.db.models import User

    code = _invite(logged_in(app))
    guest = client_for(app)
    for username, password, confirm, message in (
        ("ANA", "clave de carla", None, "Ya existe el usuario «ana»"),
        ("demo", "clave de carla", None, "reservado"),
        ("carla", "corta", None, "al menos 8"),
        ("carla", "barcelona", None, "demasiado común"),
        ("carla", "clave de carla", "otra", "no coinciden"),
    ):
        response = _register(guest, code, username, password, confirm)
        assert response.status_code == 200 and message in response.text, message
        assert "clave de carla" not in response.text                     # passwords never echoed
    assert _register(guest, code).status_code == 303                       # the code was still valid
    db = factory()
    try:
        assert db.query(User).count() == 3
    finally:
        db.close()


def test_wrong_codes_are_generic_and_throttled(app, family):
    code = _invite(logged_in(app))          # before the guessing: the IP is throttled for the login too
    guest = client_for(app)
    statuses = []
    for i in range(12):
        response = guest.post("/registro/codigo", data={"codigo": f"ZZZZ-ZZZ{i % 8 + 2}"}, headers=ORIGIN)
        statuses.append(response.status_code)
        if response.status_code == 400:
            assert "Código no válido o caducado" in response.text
    assert statuses[:4] == [400] * 4 and 429 in statuses
    # Even a valid code waits now (the throttle comes first), and so do the hint and the login.
    assert login(client_for(app), "ana", "clave-ana").status_code == 429
    assert guest.post("/registro/codigo", data={"codigo": code}, headers=ORIGIN).status_code == 429
    assert guest.get("/registro/disponible", params={"username": "x", "codigo": code}).status_code == 429


def test_name_hint_needs_a_valid_code(app, family):
    guest = client_for(app)
    response = guest.get("/registro/disponible", params={"username": "ana", "codigo": "AAAA-AAAA"})
    assert response.status_code == 403 and "ana" not in response.text


def test_two_people_with_the_same_code_create_one_account(app, family, factory, monkeypatch):
    from app.db.models import User
    from app.web.routers import auth as auth_router

    code = _invite(logged_in(app))
    first, second = client_for(app), client_for(app)
    assert first.post("/registro/codigo", data={"codigo": code}, headers=ORIGIN).status_code == 200
    assert second.post("/registro/codigo", data={"codigo": code}, headers=ORIGIN).status_code == 200
    assert _register(first, code, "carla").status_code == 303
    late = _register(second, code, "dani", "clave de dani")
    assert late.status_code == 400 and "no válido o caducado" in late.text

    # A database refusal after consume() gives the invitation back.
    code = _invite(logged_in(app))
    real = auth_router.create_user

    def racing(db, username, password):
        real(db, "Eva", "clave de eva larga")                 # someone took the name meanwhile
        return real(db, username, password)

    monkeypatch.setattr(auth_router, "create_user", racing)
    response = _register(client_for(app), code, "Eva", "clave de otra eva")
    assert response.status_code == 200 and "existe el usuario «Eva»" in response.text
    assert app.state.invites.peek(code) is not None
    db = factory()
    try:
        assert db.query(User).filter(User.username == "dani").count() == 0
    finally:
        db.close()


def test_invitations_need_a_family_member_and_csrf(app, family, factory):
    from app.config import DEMO_USERNAME

    anonymous = client_for(app)
    assert anonymous.post("/mas/invitaciones", headers=ORIGIN, follow_redirects=False).headers["location"] == "/login"
    ana = logged_in(app)
    assert ana.post("/mas/invitaciones", headers=ORIGIN).status_code == 403                  # no CSRF
    db = factory()
    create_user(db, DEMO_USERNAME, "123456", enforce_policy=False)
    db.close()
    visitor = logged_in(app, DEMO_USERNAME, "123456")
    page = visitor.get("/mas").text
    assert "/mas/invitaciones" not in page and ">beto<" not in page and ">ana<" not in page
    assert visitor.post("/mas/invitaciones", headers=_hx(visitor)).status_code == 403
    assert visitor.post("/mas/invitaciones/x/anular", headers=_hx(visitor)).status_code == 403
    assert "demostración" in ana.get("/mas").text


def test_three_invitations_at_most_and_only_my_own_can_be_cancelled(app, family):
    ana, beto = logged_in(app), logged_in(app, "beto", "clave-beto")
    for _ in range(3):
        _invite(ana)
    limited = ana.post("/mas/invitaciones", headers=_hx(ana))
    assert limited.status_code == 200 and "Ya tienes 3 invitaciones activas" in limited.text
    ids = [inv.id for inv in app.state.invites.active_by(family["ana"])]
    assert beto.post(f"/mas/invitaciones/{ids[0]}/anular", headers=_hx(beto)).status_code == 404
    response = ana.post(f"/mas/invitaciones/{ids[0]}/anular", headers=_hx(ana))
    assert response.status_code == 204 and "Invitaci" in response.headers["hx-trigger"]
    assert len(app.state.invites.active_by(family["ana"])) == 2
    _invite(ana)                                                             # room again


def test_deleting_an_account_cancels_its_invitations(app, family):
    ana = logged_in(app)
    code = _invite(ana)
    response = ana.post("/mas/cuenta/borrar", data={"confirm_name": "ana", "password": "clave-ana"},
                        headers=_hx(ana))
    assert response.status_code == 204
    assert app.state.invites.peek(code) is None


def test_logout_goes_back_to_the_login(app, family):
    client = logged_in(app)
    response = client.post("/logout", data={"csrf_token": csrf_of(client)}, headers=ORIGIN, follow_redirects=False)
    assert response.headers["location"] == "/login"


def test_invite_store_single_use_expiry_and_threads():
    import threading

    from app.web.invites import ALPHABET, InviteStore, normalize_code

    now = [1000.0]
    store = InviteStore(clock=lambda: now[0])
    code, invite = store.issue(1, "ana")
    assert len(normalize_code(code)) == 8 and set(normalize_code(code)) <= set(ALPHABET)
    assert code.replace("-", "") not in repr(store.__dict__)              # only the hash is kept
    assert store.peek(code.lower()) is invite and store.peek(code.replace("-", " ")) is invite
    now[0] += 15 * 60
    assert store.peek(code) is None and store.consume(code) is None       # expired

    code, invite = store.issue(1, "ana")
    results = []
    barrier = threading.Barrier(8)

    def grab():
        barrier.wait()
        results.append(store.consume(code))

    threads = [threading.Thread(target=grab) for _ in range(8)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert sum(r is not None for r in results) == 1
    store.restore(invite)
    assert store.peek(code) is invite
    assert store.revoke_by(1) == 1 and store.peek(code) is None
    store.issue(2, "beto")
    assert store.revoke_all() == 1


def test_switching_data_folder_or_certificates_drops_invitations(monkeypatch):
    from app.web import server as server_module

    web = server_module.WebServer()
    web.invites.issue(1, "ana")
    web.sessions.create(1, "ana", "192.168.1.2", "iPhone")
    monkeypatch.setattr(server_module, "ensure_tls", lambda regenerate=False: None)
    web.regenerate_certificates()
    assert web.invites.active_by(1) == [] and web.sessions.active() == []


# ---------------------------------------------------------------------------
# Ajustes › Mi cuenta (W10): password, name, delete — own account only
# ---------------------------------------------------------------------------

ACCOUNT_FORMS = ("/mas/cuenta/contrasena", "/mas/cuenta/nombre", "/mas/cuenta/borrar", "/mas/cuenta/datos")


def _sid(client: TestClient) -> str:
    return client.cookies.get(SESSION_COOKIE)


def test_account_routes_need_a_session_csrf_and_a_family_member(app, family, factory):
    from app.config import DEMO_USERNAME

    anonymous = client_for(app)
    for path in ACCOUNT_FORMS:
        assert anonymous.get(path, follow_redirects=False).headers["location"] == "/login"
        assert anonymous.post(path, data={"password": "x"}, headers=ORIGIN,
                              follow_redirects=False).headers["location"] == "/login"

    ana = logged_in(app)
    assert "Mi cuenta" in ana.get("/mas").text
    for path in ACCOUNT_FORMS:
        assert ana.get(path).status_code == 200
        assert ana.post(path, data={"password": "clave-ana"}, headers=ORIGIN).status_code == 403       # no CSRF
        assert ana.post(path, data={"password": "clave-ana"},
                        headers={"X-CSRF-Token": csrf_of(ana)}).status_code == 403                   # no origin

    db = factory()
    create_user(db, DEMO_USERNAME, "123456", enforce_policy=False)
    db.close()
    visitor = logged_in(app, DEMO_USERNAME, "123456")
    assert "/mas/cuenta/" not in visitor.get("/mas").text
    for path in ACCOUNT_FORMS:
        assert visitor.get(path).status_code == 403
        assert visitor.post(path, data={"password": "123456", "confirm_name": DEMO_USERNAME},
                            headers=_hx(visitor)).status_code == 403
    assert app.state.backups == []


def test_password_change_rotates_this_session_and_closes_the_others(app, family, factory):
    phone, tablet = logged_in(app), logged_in(app)
    old_sid, tablet_sid = _sid(phone), _sid(tablet)
    old_csrf = csrf_of(phone)

    response = phone.post("/mas/cuenta/contrasena", data={
        "current": "clave-ana", "password": "huerto de limones", "confirm": "huerto de limones",
    }, headers={**ORIGIN, "X-CSRF-Token": old_csrf, "HX-Request": "true"})

    assert response.status_code == 204
    trigger = response.headers["hx-trigger"]
    assert "closeSheet" in trigger and "1 dispositivo" in trigger and "csrfToken" in trigger
    sessions = app.state.sessions
    assert sessions.get(old_sid) is None and sessions.get(tablet_sid) is None
    new_sid = _sid(phone)
    assert new_sid and new_sid != old_sid and sessions.get(new_sid).csrf != old_csrf
    assert phone.get("/mas", follow_redirects=False).status_code == 200            # this device stays in
    assert tablet.get("/mas", follow_redirects=False).headers["location"] == "/login"
    assert login(client_for(app), "ana", "clave-ana").status_code == 401
    assert login(client_for(app), "ana", "huerto de limones").status_code == 303
    assert (family["ana"], "account") in app.state.changes
    db = factory()
    try:
        entry = db.query(AuditLog).filter(AuditLog.table_affected == "users").one()
        assert entry.user_id == family["ana"] and "contraseña cambiada" in entry.new_value
        assert "huerto" not in entry.new_value
    finally:
        db.close()


@pytest.mark.parametrize("data, field_message", [
    ({"current": "mal", "password": "huerto de limones", "confirm": "huerto de limones"}, "no es correcta"),
    ({"current": "clave-ana", "password": "corta", "confirm": "corta"}, "al menos 8"),
    ({"current": "clave-ana", "password": "barcelona", "confirm": "barcelona"}, "demasiado común"),
    ({"current": "clave-ana", "password": "huerto de limones", "confirm": "huerto"}, "no coinciden"),
])
def test_password_change_errors_stay_in_the_sheet(app, family, data, field_message):
    client = logged_in(app)
    sid = _sid(client)
    response = client.post("/mas/cuenta/contrasena", data=data, headers=_hx(client))
    assert response.status_code == 200 and field_message in response.text
    assert 'class="field invalid"' in response.text
    assert app.state.sessions.get(sid) is not None
    assert login(client_for(app), "ana", "clave-ana").status_code == 303


def test_wrong_passwords_in_mi_cuenta_hit_the_login_throttle(app, family):
    client = logged_in(app)
    headers = _hx(client)
    statuses = [
        client.post("/mas/cuenta/contrasena", data={"current": f"mal-{i}", "password": "huerto de limones",
                                                    "confirm": "huerto de limones"}, headers=headers).status_code
        for i in range(6)
    ]
    assert statuses[:4] == [200] * 4 and 429 in statuses[4:]
    blocked = client.post("/mas/cuenta/nombre", data={"username": "Ana María", "password": "clave-ana"},
                          headers=headers)
    assert blocked.status_code == 429 and "Demasiados intentos" in blocked.text
    # The same keys as the login: that user name waits there too.
    assert login(client_for(app), "ana", "clave-ana").status_code == 429


def test_my_data_download_needs_the_password_and_holds_only_my_rows(app, family, factory):
    import io
    import zipfile

    client = logged_in(app)
    wrong = client.post("/mas/cuenta/datos", data={"password": "mal"}, headers=_hx(client))
    assert wrong.status_code == 200 and "no es correcta" in wrong.text and "/mas/cuenta/datos/" not in wrong.text

    ready = client.post("/mas/cuenta/datos", data={"password": "clave-ana"}, headers=_hx(client))
    assert ready.status_code == 200 and "Tus datos están listos" in ready.text
    link = re.search(r'href="(/mas/cuenta/datos/[^"]+)"', ready.text).group(1)

    other = logged_in(app)                              # another session of ana: not its link
    assert other.get(link).status_code == 404
    download = client.get(link)
    assert download.status_code == 200 and download.headers["content-type"] == "application/zip"
    assert download.headers["cache-control"] == "no-store"
    assert 'attachment; filename="familyfin-datos-ana-' in download.headers["content-disposition"]
    assert client.get(link).status_code == 404          # one use only

    archive = zipfile.ZipFile(io.BytesIO(download.content))
    everything = b"".join(archive.read(n) for n in archive.namelist())
    xlsx = next(n for n in archive.namelist() if n.endswith(".xlsx"))
    from openpyxl import load_workbook
    cells = " ".join(str(c) for ws in load_workbook(io.BytesIO(archive.read(xlsx)))
                     for r in ws.iter_rows(values_only=True) for c in r if c)
    assert "MERCADO ANA" in cells and "MERCADO BETO" not in cells and "Acme" not in cells
    assert b"MERCADO BETO" not in everything

    db = factory()
    try:
        entry = db.query(AuditLog).filter(AuditLog.table_affected == "users").one()
        assert entry.user_id == family["ana"] and "descargó sus datos: 2 movimientos" in entry.new_value
    finally:
        db.close()


def test_my_data_download_hits_the_login_throttle(app, family):
    client = logged_in(app)
    statuses = [client.post("/mas/cuenta/datos", data={"password": f"mal-{i}"}, headers=_hx(client)).status_code
                for i in range(6)]
    assert statuses[:4] == [200] * 4 and 429 in statuses[4:]
    blocked = client.post("/mas/cuenta/datos", data={"password": "clave-ana"}, headers=_hx(client))
    assert blocked.status_code == 429 and "/mas/cuenta/datos/" not in blocked.text


def test_my_data_link_expires(app, family, monkeypatch):
    from app.web.routers import account as account_router

    client = logged_in(app)
    ready = client.post("/mas/cuenta/datos", data={"password": "clave-ana"}, headers=_hx(client))
    link = re.search(r'href="(/mas/cuenta/datos/[^"]+)"', ready.text).group(1)
    monkeypatch.setattr(account_router, "EXPORT_TTL", -1)
    assert client.get(link).status_code == 404


def test_rename_shows_the_new_name_in_every_open_session(app, family, factory):
    phone, tablet = logged_in(app), logged_in(app)
    taken = phone.get("/mas/cuenta/nombre/disponible", params={"username": "BETO"})
    assert "Ya existe el usuario «beto»" in taken.text
    assert "Es tu nombre actual" in phone.get("/mas/cuenta/nombre/disponible", params={"username": "ana"}).text
    assert "está libre" in phone.get("/mas/cuenta/nombre/disponible", params={"username": "Ana María"}).text

    wrong = phone.post("/mas/cuenta/nombre", data={"username": "Ana María", "password": "mal"}, headers=_hx(phone))
    assert wrong.status_code == 200 and "no es correcta" in wrong.text

    response = phone.post("/mas/cuenta/nombre", data={"username": " Ana María ", "password": "clave-ana"},
                          headers=_hx(phone))
    assert response.status_code == 204 and "refreshPage" in response.headers["hx-trigger"]
    assert "<b>Ana María</b>" in tablet.get("/mas").text
    assert login(client_for(app), "Ana María", "clave-ana").status_code == 303
    assert (family["ana"], "account") in app.state.changes
    db = factory()
    try:
        entry = db.query(AuditLog).filter(AuditLog.table_affected == "users").one()
        assert entry.old_value == "«ana»" and entry.new_value.startswith("«Ana María»")
    finally:
        db.close()


def test_delete_my_account_backs_up_first_and_signs_out_everywhere(app, family, factory):
    from app.db.models import User

    phone, tablet = logged_in(app), logged_in(app)
    sheet = phone.get("/mas/cuenta/borrar").text
    assert ">2</span> movimientos" in sheet and "Se borrarán tu usuario" in sheet
    assert 'data-delete-submit disabled' in sheet

    wrong = phone.post("/mas/cuenta/borrar", data={"confirm_name": "Ana", "password": "clave-ana"},
                       headers=_hx(phone))
    assert wrong.status_code == 200 and "no coincide" in wrong.text and app.state.backups == []

    response = phone.post("/mas/cuenta/borrar", data={"confirm_name": "ana", "password": "clave-ana"},
                          headers=_hx(phone))
    assert response.status_code == 204
    assert response.headers["hx-redirect"] == "/login?borrada=1"
    assert app.state.backups == ["ana"]
    assert not [s for s in app.state.sessions.active() if s.user_id == family["ana"]]
    assert tablet.get("/mas", follow_redirects=False).headers["location"] == "/login"
    assert (family["ana"], "account_deleted") in app.state.changes
    assert "Cuenta borrada · copia guardada en el PC" in client_for(app).get("/login?borrada=1").text

    db = factory()
    try:
        assert db.get(User, family["ana"]) is None
        assert db.query(Transaction).filter(Transaction.user_id == family["ana"]).count() == 0
        assert db.query(Transaction).filter(Transaction.user_id == family["beto"]).count() == 2
        entry = db.query(AuditLog).filter(AuditLog.table_affected == "users").one()
        assert entry.user_id is None and entry.record_id == family["ana"]
        assert "profin_predelete_ana.zip" in entry.new_value
    finally:
        db.close()
    assert login(client_for(app), "beto", "clave-beto").status_code == 303


def test_delete_is_refused_while_family_rules_depend_on_my_categories(app, family, factory):
    from app.db.models import User

    _add_web_rule(factory, family["ana"], family["ana_cat"], "REGLA DE TODOS", global_=True)
    client = logged_in(app)
    sheet = client.get("/mas/cuenta/borrar").text
    assert "1 reglas globales" in sheet and "data-delete-form" not in sheet
    response = client.post("/mas/cuenta/borrar", data={"confirm_name": "ana", "password": "clave-ana"},
                           headers=_hx(client))
    assert response.status_code == 200 and "reglas globales" in response.text
    assert app.state.backups == []
    db = factory()
    try:
        assert db.get(User, family["ana"]) is not None
    finally:
        db.close()


def test_delete_without_a_backup_deletes_nothing(app, family, factory):
    from app.db.models import User

    def broken(_name):
        raise OSError("sin espacio")

    app.state.backup_before_delete = broken
    client = logged_in(app)
    response = client.post("/mas/cuenta/borrar", data={"confirm_name": "ana", "password": "clave-ana"},
                           headers=_hx(client))
    assert response.status_code == 200 and "no se ha borrado nada" in response.text
    app.state.backup_before_delete = None
    response = client.post("/mas/cuenta/borrar", data={"confirm_name": "ana", "password": "clave-ana"},
                           headers=_hx(client))
    assert response.status_code == 503
    db = factory()
    try:
        assert db.get(User, family["ana"]) is not None
    finally:
        db.close()


def test_session_store_rotate_and_rename():
    store = SessionStore()
    a = store.create(1, "ana", "192.168.1.2", "iPhone")
    b = store.create(1, "ana", "192.168.1.3", "iPad")
    other = store.create(2, "beto", "192.168.1.4", "iPhone")
    assert store.rename_user(1, "Ana María") == 2 and store.get(b.sid).username == "Ana María"
    fresh = store.rotate(a.sid)
    assert store.get(a.sid) is None and fresh.user_id == 1 and fresh.csrf != a.csrf
    assert store.revoke_user(1, keep=fresh.sid) == 1
    assert store.get(fresh.sid) is not None and store.get(other.sid) is not None
    assert store.rotate("no-existe") is None


def test_rule_form_saves_the_movement_type_and_filters_by_it(app, family, factory):
    client = logged_in(app)
    headers = _hx(client)
    assert 'name="transaction_type"' in client.get("/mas/reglas/nueva").text
    response = client.post("/mas/reglas", data={"category_id": family["ana_cat"], "pattern": "BIZUM",
                                                "match_type": "contains", "transaction_type": "expense"},
                           headers=headers)
    assert response.status_code == 204
    response = client.post("/mas/reglas", data={"category_id": family["ana_cat"], "pattern": "X",
                                                "transaction_type": "investment"}, headers=headers)
    assert "Tipo de movimiento no válido" in response.text
    db = factory()
    try:
        rule = db.query(AutoCategorizationRule).filter_by(user_id=family["ana"]).one()
        assert rule.transaction_type == TransactionType.expense
        rule_id = rule.id
    finally:
        db.close()
    assert 'value="expense" checked' in client.get(f"/mas/reglas/{rule_id}").text

    # The ana fixture movement «MERCADO ANA» is an expense: an income-only
    # preview finds nothing.
    preview = client.get("/mas/reglas/probar", params={"pattern": "MERCADO", "match_type": "contains",
                                                       "transaction_type": "income"}).text
    assert "MERCADO ANA" not in preview
    preview = client.get("/mas/reglas/probar", params={"pattern": "MERCADO", "match_type": "contains",
                                                       "transaction_type": "expense"}).text
    assert "MERCADO ANA" in preview

    assert "«BIZUM»" in client.get("/mas/reglas", params={"movimiento": "expense"}).text
    assert "«BIZUM»" not in client.get("/mas/reglas", params={"movimiento": "any"}).text

    # Saving the form without a type sets it back to any.
    response = client.post("/mas/reglas", data={"rule_id": rule_id, "category_id": family["ana_cat"],
                                                "pattern": "BIZUM", "match_type": "contains",
                                                "transaction_type": ""}, headers=headers)
    assert response.status_code == 204
    db = factory()
    try:
        assert db.get(AutoCategorizationRule, rule_id).transaction_type is None
    finally:
        db.close()


# ---------------------------------------------------------------------------
# Fijo / traspaso marks
# ---------------------------------------------------------------------------

def test_category_nature_saves_on_tap_with_audit(app, family, factory):
    client = logged_in(app)
    headers = {**ORIGIN, "X-CSRF-Token": csrf_of(client), "HX-Request": "true"}
    url = f"/mas/categorias/{family['ana_cat']}/tipo"

    response = client.post(url, data={"nature": "traspaso"}, headers=headers)
    assert response.status_code == 200
    assert "beto" in response.text                  # the «Con» picker appears for a traspaso
    response = client.post(url, data={"nature": "traspaso", "peer": family["beto"]}, headers=headers)
    assert response.status_code == 200
    db = factory()
    try:
        cat = db.get(Category, family["ana_cat"])
        assert (cat.nature, cat.transfer_peer_user_id) == ("traspaso", family["beto"])
        assert db.query(AuditLog).filter_by(table_affected="categories", record_id=cat.id).count() == 2
    finally:
        db.close()

    assert client.post(f"/mas/categorias/{family['beto_cat']}/tipo", data={"nature": "fijo"},
                       headers=headers).status_code == 404
    bad = client.post(url, data={"nature": "mensual"}, headers=headers)
    assert "toastError" in bad.headers.get("HX-Trigger", "")
    assert "Traspaso" in client.get("/mas/categorias").text      # badge in the list


def test_movement_fixed_override_saves_and_filters(app, family, factory):
    client = logged_in(app)
    headers = {**ORIGIN, "X-CSRF-Token": csrf_of(client), "HX-Request": "true"}
    txn = family["ana_txn"]

    sheet = client.get(f"/movimientos/{txn}").text
    assert "Según categoría" in sheet
    response = client.post(f"/movimientos/{txn}/fijo", data={"fijo": "fijo"}, headers=headers)
    assert response.status_code == 200
    assert "Fijo" in response.text                  # the row's badge, swapped out of band
    assert f"tx-{txn}" in client.get("/movimientos?fijo=fijo").text
    assert f"tx-{txn}" not in client.get("/movimientos?fijo=variable").text

    assert client.post(f"/movimientos/{family['beto_txn']}/fijo", data={"fijo": "fijo"},
                       headers=headers).status_code == 404
    assert client.post(f"/movimientos/{txn}/fijo", data={"fijo": "fijo"},
                       headers=ORIGIN).status_code == 403      # no CSRF token
    db = factory()
    try:
        assert db.get(Transaction, txn).fixed_override == "fijo"
        assert db.get(Transaction, family["beto_txn"]).fixed_override is None
        assert db.query(AuditLog).filter_by(table_affected="transactions", record_id=txn).count() == 1
    finally:
        db.close()


# ---------------------------------------------------------------------------
# Previsión
# ---------------------------------------------------------------------------

def _months_of_bills(factory, user_id, cat_id, n=6, amount=-80.0, issuer="IBERDROLA"):
    """n monthly charges ending last month (enough complete months)."""
    db = factory()
    try:
        account = db.query(Account).filter_by(user_id=user_id).first()
        y, m = date.today().year, date.today().month
        for _ in range(n):
            y, m = (y - 1, 12) if m == 1 else (y, m - 1)
            db.add(Transaction(user_id=user_id, account_id=account.id, transaction_date=date(y, m, 10),
                               amount=amount, type=TransactionType.expense, issuer_name=issuer,
                               category_id=cat_id))
        db.commit()
    finally:
        db.close()


def test_prevision_page_mine_and_family(app, family, factory):
    db = factory()
    try:
        luz = Category(user_id=family["ana"], parent_id=db.get(Category, family["ana_cat"]).parent_id,
                       name="Luz", level=CategoryLevel.category, nature="fijo")
        db.add(luz)
        db.commit()
        luz_id = luz.id
    finally:
        db.close()
    _months_of_bills(factory, family["ana"], luz_id)

    client = logged_in(app)
    page = client.get("/prevision")
    assert page.status_code == 200
    assert "Gastos previstos" in page.text and "IBERDROLA" in page.text
    assert "80,00" in page.text
    family_page = client.get("/prevision?vista=familia")
    assert family_page.status_code == 200
    assert "ana, beto" in family_page.text
    assert "MERCADO BETO" not in family_page.text               # never another member's movements
    data = client.get("/api/prevision").json()
    assert data["fixed_baseline"] == 80.0 and data["months"] == 6


def test_prevision_family_view_is_closed_to_demo(app, factory):
    db = factory()
    try:
        create_user(db, "DEMO", "clave-demo", enforce_policy=False)
    finally:
        db.close()
    client = logged_in(app, "DEMO", "clave-demo")
    page = client.get("/prevision")
    assert page.status_code == 200 and "Familia" not in page.text
    assert client.get("/prevision?vista=familia").status_code == 403
    assert client.get("/api/prevision?vista=familia").status_code == 403


def test_prevision_suggestions_mark_categories(app, family, factory):
    db = factory()
    try:
        gasto_id = db.get(Category, family["ana_cat"]).parent_id
        hipoteca = Category(user_id=family["ana"], parent_id=gasto_id, name="Hipoteca",
                            level=CategoryLevel.category)
        db.add(hipoteca)
        db.commit()
        hipoteca_id = hipoteca.id
    finally:
        db.close()
    client = logged_in(app)
    page = client.get("/prevision").text
    assert "Marca tus gastos fijos" in page and f'name="cat_{hipoteca_id}"' in page
    headers = {**ORIGIN, "X-CSRF-Token": csrf_of(client), "HX-Request": "true"}
    response = client.post("/prevision/marcar", data={f"cat_{hipoteca_id}": "fijo",
                                                      f"cat_{family['beto_cat']}": "fijo"}, headers=headers)
    assert response.status_code == 404                           # another member's category
    response = client.post("/prevision/marcar", data={f"cat_{hipoteca_id}": "fijo"}, headers=headers)
    assert response.status_code == 204
    db = factory()
    try:
        assert db.get(Category, hipoteca_id).nature == "fijo"
        assert db.get(Category, family["beto_cat"]).nature is None
    finally:
        db.close()
    assert "Marca tus gastos fijos" not in client.get("/prevision").text


def test_prevision_lists_income_apart_from_expenses(app, family, factory):
    db = factory()
    try:
        gasto_id = db.get(Category, family["ana_cat"]).parent_id
        ingreso = Category(user_id=family["ana"], name="INGRESO", level=CategoryLevel.type,
                           transaction_type=TransactionType.income)
        db.add(ingreso)
        db.flush()
        nomina = Category(user_id=family["ana"], parent_id=ingreso.id, name="Nómina",
                          level=CategoryLevel.category, nature="fijo")
        luz = Category(user_id=family["ana"], parent_id=gasto_id, name="Luz",
                       level=CategoryLevel.category, nature="fijo")
        db.add_all([nomina, luz])
        db.commit()
        nomina_id, luz_id = nomina.id, luz.id
    finally:
        db.close()
    _months_of_bills(factory, family["ana"], luz_id, amount=-80.0, issuer="IBERDROLA")
    _months_of_bills(factory, family["ana"], nomina_id, amount=2100.0, issuer="EMPRESA SA")
    db = factory()
    try:
        account = db.query(Account).filter_by(user_id=family["ana"]).first()
        for cat_id, amount, issuer in ((luz_id, -80.0, "IBERDROLA"), (nomina_id, 2100.0, "EMPRESA SA")):
            db.add(Transaction(user_id=family["ana"], account_id=account.id, transaction_date=date.today(),
                               amount=amount, type=TransactionType.expense, issuer_name=issuer,
                               category_id=cat_id))       # this month's, so both are still active
        db.flush()
        for t in db.query(Transaction).filter_by(issuer_name="EMPRESA SA"):
            t.type = TransactionType.income                # the salary rows are income
        db.commit()
    finally:
        db.close()

    page = logged_in(app).get("/prevision").text
    incomes = page[page.index('aria-label="Ingresos previstos"'):page.index("Total ingresos previstos")]
    expenses = page[page.index('aria-label="Gastos previstos"'):page.index("Total gastos previstos")]
    assert "EMPRESA SA" in incomes and "EMPRESA SA" not in expenses
    assert "IBERDROLA" in expenses and "IBERDROLA" not in incomes
    assert "+2.100,00" in incomes
    data = logged_in(app).get("/api/prevision").json()["next_month"]
    assert [i["label"] for i in data["income"]["items"]] == ["EMPRESA SA"]
    assert [i["label"] for i in data["expenses"]["items"]] == ["IBERDROLA"]


# ---------------------------------------------------------------------------
# Resumen: custom range
# ---------------------------------------------------------------------------

def test_resumen_custom_range(app, family, factory):
    db = factory()
    try:
        account = db.query(Account).filter_by(user_id=family["ana"]).first()
        db.add(Transaction(user_id=family["ana"], account_id=account.id, transaction_date=date(2025, 3, 15),
                           amount=-300.0, type=TransactionType.expense, issuer_name="VIAJE MARZO"))
        db.commit()
    finally:
        db.close()
    client = logged_in(app)

    page = client.get("/?rango=custom&desde=2025-03-01&hasta=2025-03-31")
    assert page.status_code == 200
    form_tag = re.search(r'<form id="custom-range"[^>]*>', page.text).group(0)
    assert "hidden" not in form_tag                                   # the dates stay visible
    assert "hidden" in re.search(r'<form id="custom-range"[^>]*>', client.get("/").text).group(0)
    assert "01/03/2025" in page.text and "31/03/2025" in page.text
    data = client.get("/api/resumen?rango=custom&desde=2025-03-31&hasta=2025-03-01").json()   # any order
    assert (data["date_from"], data["date_to"]) == ("2025-03-01", "2025-03-31")
    assert data["expense"] == 300.0 and data["income"] == 0
    # Without both dates it falls back to the month.
    month = client.get("/api/resumen?rango=custom&desde=2025-03-01").json()
    assert month["date_from"] == date.today().replace(day=1).isoformat()
    assert month["expense"] == 42.5
    assert "Personalizado" in client.get("/").text
